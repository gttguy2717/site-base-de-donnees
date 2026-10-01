import sys
from docx import Document

doc = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

print("=== SOMMAIRE (FR) ===")
for i in range(29, 53):
    print(f"[{i}] {doc.paragraphs[i].text.strip()}")

print("\n=== LISTE DES FIGURES (FR) ===")
for i in range(54, 85):
    print(f"[{i}] {doc.paragraphs[i].text.strip()}")

print("\n=== LISTE DES TABLEAUX (FR) ===")
for i in range(86, 92):
    print(f"[{i}] {doc.paragraphs[i].text.strip()}")
