import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc_en = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_ENGLISH.docx')
doc_fr = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

# Let's inspect P[105] in FR and its aligned paragraph in EN
# In FR, P[105] is "Le projet a pour thème..."
# In EN, what index is "The project's theme is..."?
for i, p in enumerate(doc_en.paragraphs):
    if "theme" in p.text.lower() and "project" in p.text.lower():
        print(f"EN P[{i}]: text={repr(p.text[:60])}")
        for j, r in enumerate(p.runs):
            print(f"  Run {j}: bold={r.bold} italic={r.italic} text={repr(r.text[:50])}")
        break
