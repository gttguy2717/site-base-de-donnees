import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc_fr = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

print(f"Total paragraphs: {len(doc_fr.paragraphs)}")
non_empty = [i for i, p in enumerate(doc_fr.paragraphs) if p.text.strip()]
print(f"Non-empty paragraphs: {len(non_empty)}")

# Check run counts
run_counts = {}
for i in non_empty:
    n = len(doc_fr.paragraphs[i].runs)
    run_counts[n] = run_counts.get(n, 0) + 1

print("Run counts distribution:", sorted(run_counts.items()))
