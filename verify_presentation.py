import os
import sys
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation("Soutenance stage plateforme numerique Soutarah Group.pptx")
print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx+1:02d} ====================")
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text.strip():
            print(f"[{s.name}] (top={round(s.top/914400,2)}, left={round(s.left/914400,2)}):")
            for p in s.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    print(f"   - {txt}")
        elif s.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print(f"[PICTURE: {s.name}] ({round(s.width/914400,2)}x{round(s.height/914400,2)} in, top={round(s.top/914400,2)}, left={round(s.left/914400,2)})")


