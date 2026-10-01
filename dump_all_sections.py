import docx
import json

doc = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

sections_def = [
    (0, 112, 'sec0_fr.json'),
    (112, 223, 'sec1_fr.json'),
    (223, 399, 'sec2_fr.json'),
    (399, 517, 'sec3_fr.json'),
    (517, 641, 'sec4_fr.json'),
    (641, len(doc.paragraphs), 'sec5_fr.json')
]

for start, end, filename in sections_def:
    items = []
    for i in range(start, min(end, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        if p.text.strip():
            items.append({
                'idx': i,
                'style': p.style.name,
                'text': p.text
            })
    with open(filename, 'w', encoding='utf-8') as f:
        json.dump(items, f, ensure_ascii=False, indent=2)
    print(f"{filename}: {len(items)} paragraphs (range {start}..{end})")

tables_data = []
for t_idx, table in enumerate(doc.tables):
    t_rows = []
    for r_idx, row in enumerate(table.rows):
        r_cells = []
        for c_idx, cell in enumerate(row.cells):
            c_paras = [p.text for p in cell.paragraphs if p.text.strip()]
            r_cells.append(c_paras)
        t_rows.append(r_cells)
    tables_data.append(t_rows)

with open('tables_fr.json', 'w', encoding='utf-8') as f:
    json.dump(tables_data, f, ensure_ascii=False, indent=2)
print("Saved tables_fr.json")
