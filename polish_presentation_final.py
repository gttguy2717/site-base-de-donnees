import os
import sys
from PIL import Image
import pptx
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

sys.stdout.reconfigure(encoding='utf-8')

# --- CHARTE OFFICIELLE DU LOGO SOUTARAH GROUP ---
LOGO_DARK_GREEN  = RGBColor(8, 44, 27)     # #082C1B - Vert forêt profond
LOGO_BRAND_GREEN = RGBColor(40, 129, 52)   # #288134 - Vert principal Soutarah
LOGO_LEAF_LIME   = RGBColor(120, 183, 43)  # #78B72B - Vert feuille vif du logo
LIME_LIGHT_BG    = RGBColor(236, 248, 226) # #ECF8E2 - Vert clair
TEXT_DARK        = RGBColor(15, 23, 42)    # #0F172A - Titre sombre
TEXT_BODY        = RGBColor(51, 65, 85)    # #334155 - Corps de texte

def place_image_proportional(slide, img_path, left, top, width, height):
    """Place une image en respectant son ratio d'aspect strict, centrée dans la zone."""
    if not os.path.exists(img_path):
        print(f"Attention: Image non trouvée -> {img_path}")
        return None
    with Image.open(img_path) as im:
        img_w, img_h = im.size
    
    img_aspect = img_w / img_h
    box_aspect = width / height

    if img_aspect > box_aspect:
        final_w = width
        final_h = width / img_aspect
        final_left = left
        final_top = top + (height - final_h) / 2
    else:
        final_h = height
        final_w = height * img_aspect
        final_top = top
        final_left = left + (width - final_w) / 2

    return slide.shapes.add_picture(img_path, final_left, final_top, width=final_w, height=final_h)

def polish():
    file_path = "Soutenance stage plateforme numerique Soutarah Group.pptx"
    print(f"Ouverture de {file_path}...")
    prs = pptx.Presentation(file_path)

    # 1. RE-NUMÉROTATION UNIFORME DE TOUS LES BADGES DE SLIDES (01 à 40)
    print("Correction des numéros de badges sur l'ensemble des 40 diapositives...")
    for idx, slide in enumerate(prs.slides):
        expected_num = f"{idx+1:02d}."
        for s in slide.shapes:
            if s.has_text_frame:
                t = s.text_frame.text.strip()
                if t.endswith('.') and len(t) <= 4 and t[:-1].isdigit():
                    s.text_frame.paragraphs[0].text = expected_num
                    s.text_frame.paragraphs[0].font.color.rgb = LOGO_LEAF_LIME
                    s.text_frame.paragraphs[0].font.bold = True

    # 2. SLIDE 01 : TITRE OFFICIEL INTÉGRAL DU RAPPORT
    print("Mise à jour Slide 01 : Titre officiel...")
    s1 = prs.slides[0]
    for s in s1.shapes:
        if s.has_text_frame:
            txt = s.text_frame.text
            if "Conception et développement" in txt:
                s.text_frame.text = "Conception et développement d’une plateforme numérique de gestion des services de SOUTARAH GROUP avec application mobile dédiée à la réservation de véhicules"
                for p in s.text_frame.paragraphs:
                    p.font.name = "Aileron Bold"
                    p.font.size = Pt(28)
                    p.font.color.rgb = TEXT_DARK
            elif "David SORHO" in txt or "Prévisualisation" in txt:
                s.text_frame.text = "David SORHO • Institut National Polytechnique Félix Houphouët-Boigny (INP-HB) • ESI / STIC"
                for p in s.text_frame.paragraphs:
                    p.font.name = "Aileron"
                    p.font.size = Pt(16)
                    p.font.color.rgb = LOGO_LEAF_LIME

    # 3. SLIDE 06 : PRÉSENTATION NETTE DE SOUTARAH GROUP SANS DOUBLONS
    print("Mise à jour Slide 06 : Contenu Soutarah Group...")
    s6 = prs.slides[5]
    for s in s6.shapes:
        if s.has_text_frame and ("Entreprise multisectorielle" in s.text_frame.text or "SOUTARAH" in s.text_frame.text):
            if s.name == "object 2":
                tf = s.text_frame
                tf.word_wrap = True
                tf.clear()
                
                # En-tête
                p0 = tf.paragraphs[0]
                p0.text = "Entreprise multisectorielle d'excellence"
                p0.font.name = "Aileron Bold"
                p0.font.size = Pt(20)
                p0.font.color.rgb = TEXT_DARK
                p0.space_after = Pt(10)
                
                # 6 Pôles
                p1 = tf.add_paragraph()
                p1.text = "• 6 Pôles d’activités : Location de véhicules, Négoce / Import-Export, Prestations techniques, Énergies renouvelables, Agropastorale et Immobilier."
                p1.font.size = Pt(14)
                p1.font.color.rgb = TEXT_BODY
                p1.space_after = Pt(8)

                # Partenaires
                p2 = tf.add_paragraph()
                p2.text = "• Partenaires stratégiques : BESSAC, DM Company, CIM IVOIRE, Southcomp Polaris, Enabel."
                p2.font.size = Pt(14)
                p2.font.color.rgb = TEXT_BODY
                p2.space_after = Pt(8)

                # Mission
                p3 = tf.add_paragraph()
                p3.text = "• Mission : Proposer des prestations fiables, innovantes et sur-mesure pour une clientèle B2B et B2C exigeante."
                p3.font.size = Pt(14)
                p3.font.color.rgb = TEXT_BODY
                p3.space_after = Pt(8)

                # Valeurs SAPE
                p4 = tf.add_paragraph()
                p4.text = "• Valeurs S.A.P.E : Solution pérenne et innovante, Adaptabilité, Priorité client, Efficacité du personnel."
                p4.font.size = Pt(14)
                p4.font.color.rgb = LOGO_BRAND_GREEN
                p4.font.bold = True

    # 4. SLIDE 17 : CAS D'UTILISATION - MISE EN VALEUR DU GRAND DIAGRAMME UML
    print("Mise à jour Slide 17 : Diagramme de cas d'utilisation grand format...")
    s17 = prs.slides[16]
    
    # Supprimer les anciennes petites images / cartes de slide 17
    shapes_to_remove = []
    for s in s17.shapes:
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
            shapes_to_remove.append(s)
        elif s.name in ["object 6", "object 7", "object 8"]:
            shapes_to_remove.append(s)
    
    for s in shapes_to_remove:
        s17.shapes._spTree.remove(s._element)
    
    # Ajouter un grand bloc de texte structuré à gauche
    left_tb = s17.shapes.add_textbox(Inches(3.32), Inches(3.48), Inches(7.2), Inches(6.0))
    tf = left_tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Vue d'ensemble des cas d'utilisation"
    p.font.name = "Aileron Bold"
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(10)
    
    p = tf.add_paragraph()
    p.text = "• Acteurs du système : Visiteur (libre consultation), Client connecté (commandes & réservations) et Administrateur (supervision globale)."
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_BODY
    p.space_after = Pt(8)

    p = tf.add_paragraph()
    p.text = "• Parcours Client : Exploration du catalogue multi-services, réservation de véhicules avec contrôle anti-conflit de dates, panier mixte, devis pro-forma PDF instantané et paiement sécurisé."
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_BODY
    p.space_after = Pt(8)

    p = tf.add_paragraph()
    p.text = "• Pilotage Administrateur : Gestion de la flotte de véhicules, mise à jour des catalogues, traitement et validation des devis, suivi des indicateurs clés (KPIs)."
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_BODY
    p.space_after = Pt(8)

    p = tf.add_paragraph()
    p.text = "• Synchronisation globale : Centralisation des règles métier via l'API REST Node.js/Express et persistance MySQL garantissant la cohérence Web-Mobile."
    p.font.size = Pt(14)
    p.font.color.rgb = LOGO_BRAND_GREEN
    p.font.bold = True

    # Insérer le diagramme de cas d'utilisation grand format à droite
    uc_img = "diagrammes/cas d'utilisation/cas_d'utilisation.png"
    if not os.path.exists(uc_img):
        uc_img = "extracted_pptx_images/slide_20_img_7_image.png"
    place_image_proportional(s17, uc_img, Inches(10.94), Inches(3.48), Inches(8.04), Inches(6.03))

    # 5. SLIDE 30 : BASE DE DONNÉES - SUPPRESSION DU TEXTE TECHNIQUE FACTICE
    print("Mise à jour Slide 30 : Suppression consigne de schéma relationnel...")
    s30 = prs.slides[29]
    for s in s30.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if "Insérer le schéma MySQL" in p.text:
                    p.text = "Structure relationnelle en 3NF sous MySQL garantissant l'intégrité référentielle, l'absence de redondance et la traçabilité complète des transactions."
                    p.font.color.rgb = LOGO_BRAND_GREEN

    # 6. SLIDE 37 : BACK-OFFICE - MISE EN VALEUR DE L'INTERFACE ADMIN RÉELLE
    print("Mise à jour Slide 37 : Dashboard Back-Office grand format...")
    s37 = prs.slides[36]
    
    # Supprimer les anciennes cartes de slide 37
    shapes_to_remove = []
    for s in s37.shapes:
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
            shapes_to_remove.append(s)
        elif s.name in ["object 6", "object 7", "object 8"]:
            shapes_to_remove.append(s)
    
    for s in shapes_to_remove:
        s37.shapes._spTree.remove(s._element)

    # Ajouter texte structuré à gauche
    left_tb = s37.shapes.add_textbox(Inches(3.32), Inches(3.48), Inches(7.2), Inches(6.0))
    tf = left_tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Supervision et administration centralisée"
    p.font.name = "Aileron Bold"
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(10)

    p = tf.add_paragraph()
    p.text = "• Tableau de bord KPIs : Vue consolidée en temps réel (chiffre d'affaires, volume de réservations, devis en attente et statut de la flotte)."
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_BODY
    p.space_after = Pt(8)

    p = tf.add_paragraph()
    p.text = "• Gestion de la flotte : Ajout, modification et suivi des véhicules, paramétrage des tarifs journaliers et historique de maintenance."
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_BODY
    p.space_after = Pt(8)

    p = tf.add_paragraph()
    p.text = "• Traitement des devis & commandes : Traçabilité de chaque demande, consultation du devis pro-forma PDF et marquage 'Traité / En cours'."
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_BODY
    p.space_after = Pt(8)

    p = tf.add_paragraph()
    p.text = "• Sécurité & Rôles : Contrôle des accès (ADMIN / CLIENT) et journalisation de toutes les actions opérationnelles."
    p.font.size = Pt(14)
    p.font.color.rgb = LOGO_BRAND_GREEN
    p.font.bold = True

    # Insérer la capture du dashboard Back-Office grand format à droite
    adm_img = "extracted_pptx_images/slide_43_img_7_image.png"
    place_image_proportional(s37, adm_img, Inches(10.94), Inches(3.48), Inches(8.04), Inches(5.8))

    # 7. SLIDE 40 : BILAN, ROI ET MÉTRIQUES RÉELLES DU RAPPORT
    print("Mise à jour Slide 40 : Métriques et ROI...")
    s40 = prs.slides[39]
    for s in s40.shapes:
        if s.has_text_frame:
            txt = s.text_frame.text
            if "Bénéfices opérationnels" in txt or "Traitement administratif" in txt:
                s.text_frame.text = "Gains Opérationnels & Financiers\n• Réduction de plus de 80% du temps de traitement des devis (< 3 min).\n• Budget total de 290 000 FCFA parfaitement maîtrisé grâce à la stack open source.\n• Zéro conflit de réservation grâce à l'algorithme anti-chevauchement."
                for p in s.text_frame.paragraphs:
                    p.font.size = Pt(13)
                    p.font.color.rgb = TEXT_BODY
            elif "Impact business" in txt or "réactivité commerciale" in txt:
                s.text_frame.text = "Impact Stratégique & ROI\n• Disponibilité 24h/24 et 7j/7 du catalogue et des devis pro-forma.\n• Rentabilité immédiate dès la première année d'exploitation.\n• Traçabilité complète des demandes et satisfaction client renforcée."
                for p in s.text_frame.paragraphs:
                    p.font.size = Pt(13)
                    p.font.color.rgb = TEXT_BODY

    # Sauvegarde finale
    prs.save(file_path)
    print(f"POLISH TERMINÉ AVEC SUCCÈS : '{file_path}' mis à jour.")

if __name__ == "__main__":
    polish()
