# -*- coding: utf-8 -*-
"""
Générateur de la présentation PowerPoint Soutarah Négoce Électrique
Basé sur le modèle et design de Présentation1.pptx (Widescreen 16:9, Dark Premium Tech)
"""

import os
import shutil
import qrcode
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ─── COULEURS DE LA CHARTE SOUTARAH ÉLECTRIQUE ───
COLOR_BG_DARK       = RGBColor(11, 17, 32)      # #0B1120 (Bleu nuit profond)
COLOR_CARD_BG       = RGBColor(26, 38, 57)      # #1A2639 (Fond des cartes)
COLOR_CARD_BORDER   = RGBColor(51, 65, 85)      # #334155 (Bordure subtile)
COLOR_ACCENT_GOLD   = RGBColor(245, 158, 11)    # #F59E0B (Or / Orange électrique)
COLOR_ACCENT_AMBER  = RGBColor(251, 191, 36)    # #FBBF24 (Jaune électrique clair)
COLOR_ACCENT_CYAN   = RGBColor(14, 165, 233)    # #0EA5E9 (Cyan technologique)
COLOR_ACCENT_GREEN  = RGBColor(16, 185, 129)    # #10B981 (Vert stock dispo)
COLOR_ACCENT_RED    = RGBColor(239, 68, 68)     # #EF4444 (Alerte bêtisier)
COLOR_TEXT_WHITE    = RGBColor(255, 255, 255)   # #FFFFFF (Blanc éclatant)
COLOR_TEXT_LIGHT    = RGBColor(226, 232, 240)   # #E2E8F0 (Gris très clair)
COLOR_TEXT_MUTED    = RGBColor(148, 163, 184)   # #94A3B8 (Gris bleuté secondaire)

FONT_TITLE = "Segoe UI"
FONT_BODY  = "Segoe UI"

LOGO_SOUTARAH = "dist/logo-soutarah.png"
LOGO_SOGELEC  = "dist/img/sogelec.jpeg"

def ensure_qr_code():
    qr_path = "qrcode_soutarah.png"
    if not os.path.exists(qr_path):
        qr = qrcode.QRCode(
            version=1,
            error_correction=qrcode.constants.ERROR_CORRECT_M,
            box_size=10,
            border=2,
        )
        qr.add_data("https://www.soutarahgroup.com")
        qr.make(fit=True)
        img = qr.make_image(fill_color="#0F172A", back_color="white")
        img.save(qr_path)
    return qr_path

def apply_slide_background(slide):
    """Crée un fond sombre et moderne 16:9."""
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_BG_DARK
    bg_shape.line.color.rgb = COLOR_BG_DARK
    bg_shape.line.width = Pt(0)
    
    # Ligne d'accent lumineuse en haut (13.333 pouces x 0.06 pouces)
    top_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), Inches(13.333), Inches(0.06)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT_GOLD
    top_bar.line.fill.background()
    return bg_shape

def add_header(slide, badge_text, title_text, subtitle_text=None, badge_color=COLOR_ACCENT_GOLD):
    """Ajoute l'en-tête standardisé avec logos, badge et titres."""
    # Logo Soutarah en haut à gauche
    if os.path.exists(LOGO_SOUTARAH):
        slide.shapes.add_picture(LOGO_SOUTARAH, Inches(0.6), Inches(0.25), width=Inches(1.2))

    # Logo SOGELEC en haut à droite
    if os.path.exists(LOGO_SOGELEC):
        slide.shapes.add_picture(LOGO_SOGELEC, Inches(11.2), Inches(0.28), width=Inches(1.5))

    # Badge de catégorie
    if badge_text:
        badge_box = slide.shapes.add_textbox(Inches(2.2), Inches(0.25), Inches(8.5), Inches(0.35))
        tf_b = badge_box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.LEFT
        run_b = p_b.add_run()
        run_b.text = badge_text.upper()
        run_b.font.name = FONT_BODY
        run_b.font.size = Pt(11)
        run_b.font.bold = True
        run_b.font.color.rgb = badge_color

    # Titre de la slide
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.7), Inches(12.133), Inches(0.65))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.alignment = PP_ALIGN.LEFT
    run_t = p_t.add_run()
    run_t.text = title_text
    run_t.font.name = FONT_TITLE
    run_t.font.size = Pt(24)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_TEXT_WHITE

    # Sous-titre optionnel
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.35), Inches(12.133), Inches(0.4))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        run_s = p_s.add_run()
        run_s.text = subtitle_text
        run_s.font.name = FONT_BODY
        run_s.font.size = Pt(13)
        run_s.font.italic = True
        run_s.font.color.rgb = COLOR_TEXT_MUTED

def add_footer(slide, slide_num):
    """Ajoute le pied de page moderne."""
    footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.133), Inches(0.35))
    tf = footer_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r1 = p.add_run()
    r1.text = "SOUTARAH GROUP × SOGELEC  •  Négoce Électrique 2.0  •  www.soutarahgroup.com"
    r1.font.name = FONT_BODY
    r1.font.size = Pt(10)
    r1.font.color.rgb = COLOR_TEXT_MUTED
    
    r2 = p.add_run()
    r2.text = f"                                                                                               {slide_num} / 6"
    r2.font.name = FONT_BODY
    r2.font.size = Pt(10)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_ACCENT_GOLD

def create_card(slide, left, top, width, height, border_color=COLOR_CARD_BORDER, bg_color=COLOR_CARD_BG):
    """Crée une carte visuelle avec coins arrondis et bordure élégante."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    return card

print("Helper functions ready.")
