import pptx
import sys
sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('Soutenance stage plateforme numerique Soutarah Group.pptx')

print(f"Total slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    texts = []
    has_img = any(s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    texts.append(t)
    # Find meaningful title
    title = "SANS TITRE"
    for t in texts:
        if t not in ['David Sorho', '9/15/2026', f'{i+1:02d}.', f'{i+1}.'] and len(t) > 3:
            title = t
            break
    print(f"Slide {i+1:02d} [{'IMG' if has_img else 'NO-IMG'}]: {title}")
