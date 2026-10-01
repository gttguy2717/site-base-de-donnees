import sys
from docx import Document

doc = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_ENGLISH.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

txbx = doc._element.xpath('.//*[local-name()="txbxContent"]')
print(f"Total textboxes: {len(txbx)}")
for i, tb in enumerate(txbx[:15]):
    t = ' '.join([x.text for x in tb.xpath('.//*[local-name()="t"]') if x.text]).strip()
    print(f"  TB[{i}]: {t}")

for i in [1, 2, 5, 6, 7, 8, 29, 30, 54, 55, 113, 114, 116, 131]:
    if i < len(doc.paragraphs):
        t = doc.paragraphs[i].text.strip()
        if t:
            print(f"  P[{i}] ({doc.paragraphs[i].style.name}): {t[:70]}")
