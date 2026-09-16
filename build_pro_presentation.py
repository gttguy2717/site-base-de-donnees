import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

# --- PALETTE DE COULEURS PROFESSIONNELLE EXÉCUTIVE ---
DARK_BG = RGBColor(6, 40, 23)           # #062817 Vert forêt profond impérial
DARK_CARD = RGBColor(12, 54, 33)         # #0C3621 Carte sur fond sombre
DARK_BORDER = RGBColor(22, 101, 52)      # #166534 Bordure carte sombre
LIGHT_BG = RGBColor(248, 250, 252)      # #F8FAFC Fond clair moderne
CARD_BG = RGBColor(255, 255, 255)        # #FFFFFF Blanc pur
CARD_BORDER = RGBColor(226, 232, 240)    # #E2E8F0 Bordure subtile
BRAND_GREEN = RGBColor(22, 163, 74)      # #16A34A Vert émeraude soutarah
EMERALD_LIGHT = RGBColor(220, 252, 231)  # #DCFCE7 Vert doux
EMERALD_DARK = RGBColor(21, 128, 61)     # #15803D Vert accent
BRAND_GOLD = RGBColor(217, 119, 6)       # #D97706 Or chaud
GOLD_LIGHT = RGBColor(254, 243, 199)     # #FEF3C7 Ambre doux
TEXT_DARK = RGBColor(15, 23, 42)         # #0F172A Slate foncé
TEXT_BODY = RGBColor(51, 65, 85)         # #334155 Slate texte
TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B Slate estompé
TEXT_WHITE = RGBColor(255, 255, 255)     # Blanc
TABLE_HEADER_BG = RGBColor(11, 61, 36)   # #0B3D24 Vert émeraude foncé pour en-tête tableau
ROW_ALT_BG = RGBColor(241, 245, 249)     # #F1F5F9 Gris bleuté très clair

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"
TOTAL_SLIDES = 50

IMG_DIR = "extracted_pptx_images"
LOGO_PATH = "public/logo-soutarah.png"

def get_img(name):
    p = os.path.join(IMG_DIR, name)
    return p if os.path.exists(p) else None

def init_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def set_slide_background(slide, prs, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, kicker, title, subtitle=None, is_dark=False):
    tb_k = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.73), Inches(0.28))
    tf_k = tb_k.text_frame
    tf_k.word_wrap = True
    tf_k.margin_left = tf_k.margin_top = tf_k.margin_right = tf_k.margin_bottom = 0
    p_k = tf_k.paragraphs[0]
    p_k.text = kicker.upper()
    p_k.font.name = FONT_HEADING
    p_k.font.size = Pt(10)
    p_k.font.bold = True
    p_k.font.color.rgb = BRAND_GOLD if is_dark else BRAND_GREEN

    tb_t = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.73), Inches(0.55))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(21)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE if is_dark else TEXT_DARK

    if subtitle:
        tb_s = slide.shapes.add_textbox(Inches(0.8), Inches(1.30), Inches(11.73), Inches(0.32))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = RGBColor(203, 213, 225) if is_dark else TEXT_MUTED

def add_footer(slide, slide_num, is_dark=False):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.73), Inches(0.012))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(20, 70, 45) if is_dark else RGBColor(226, 232, 240)
    line.line.fill.background()

    tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(8.5), Inches(0.25))
    tf_l = tb_l.text_frame
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "SOUTARAH GROUP • Soutenance de Stage • David SORHO • INP-HB / ESI"
    p_l.font.name = FONT_BODY
    p_l.font.size = Pt(9)
    p_l.font.color.rgb = RGBColor(148, 163, 184) if is_dark else TEXT_MUTED

    tb_r = slide.shapes.add_textbox(Inches(10.53), Inches(7.12), Inches(2.0), Inches(0.25))
    tf_r = tb_r.text_frame
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.text = f"{slide_num:02d} / {TOTAL_SLIDES:02d}"
    p_r.font.name = FONT_HEADING
    p_r.font.size = Pt(9.5)
    p_r.font.bold = True
    p_r.font.color.rgb = BRAND_GOLD if is_dark else BRAND_GREEN

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    
    if top_accent_color:
        accent = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.08))
        accent.fill.solid()
        accent.fill.fore_color.rgb = top_accent_color
        accent.line.fill.background()
    return card

def add_badge(slide, left, top, width, height, text, bg_color=EMERALD_LIGHT, txt_color=EMERALD_DARK):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    pill.fill.solid()
    pill.fill.fore_color.rgb = bg_color
    pill.line.fill.background()
    tf = pill.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text.upper()
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_HEADING
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = txt_color
    return pill

def create_section_divider(prs, slide_num, seq_num, part_title, part_desc):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    set_slide_background(slide, prs, DARK_BG)

    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.8))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = BRAND_GOLD
    accent_bar.line.fill.background()

    add_badge(slide, Inches(1.2), Inches(1.8), Inches(1.8), Inches(0.35), f"SÉQUENCE {seq_num:02d}", bg_color=DARK_CARD, txt_color=BRAND_GOLD)

    tb_t = slide.shapes.add_textbox(Inches(1.2), Inches(2.35), Inches(11.0), Inches(1.5))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = part_title
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(32)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE

    tb_d = slide.shapes.add_textbox(Inches(1.2), Inches(4.0), Inches(10.5), Inches(1.2))
    tf_d = tb_d.text_frame
    tf_d.word_wrap = True
    tf_d.margin_left = tf_d.margin_top = tf_d.margin_right = tf_d.margin_bottom = 0
    p_d = tf_d.paragraphs[0]
    p_d.text = part_desc
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(14)
    p_d.font.color.rgb = RGBColor(203, 213, 225)

    add_footer(slide, slide_num, is_dark=True)
    return slide

print("Framework setup ready.")
