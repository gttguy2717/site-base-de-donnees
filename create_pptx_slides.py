import os
from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]

# Slide 1: Grid 3x2 (Cards)
slide1 = prs.slides.add_slide(blank_slide_layout)
img1_path = "presentation_services_soutarah_grid3x2.png"
if os.path.exists(img1_path):
    slide1.shapes.add_picture(img1_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

# Slide 2: Dual 2x3 (Horizontal Cards)
slide2 = prs.slides.add_slide(blank_slide_layout)
img2_path = "presentation_services_soutarah_dual2x3.png"
if os.path.exists(img2_path):
    slide2.shapes.add_picture(img2_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

output_pptx = "Presentation_Services_Soutarah.pptx"
prs.save(output_pptx)
print(f"Successfully saved {output_pptx} with {len(prs.slides)} slides!")
