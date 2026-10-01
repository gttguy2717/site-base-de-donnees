import docx
import json
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc_fr = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')
doc_en = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_ENGLISH.docx')

with open('fr_clean_474.json', encoding='utf-8') as f:
    fr_paras = json.load(f)

en_paras = [(i, p) for i, p in enumerate(doc_en.paragraphs) if p.text.strip()]

# Step through alignment
fr_idx = 0
en_idx = 0
pairs = []
while fr_idx < len(fr_paras) and en_idx < len(en_paras):
    f_item = fr_paras[fr_idx]
    e_idx, e_p = en_paras[en_idx]
    f_txt = f_item['text'].strip()
    if 'genius pay' in f_txt.lower() or 'passerelle de paiement' in f_txt.lower():
        fr_idx += 1
        continue
    f_p = doc_fr.paragraphs[f_item['idx']]
    pairs.append((f_item['idx'], f_p, e_idx, e_p))
    fr_idx += 1
    en_idx += 1

print(f"Total aligned pairs: {len(pairs)}")

# Compare run counts
same_runs = 0
diff_runs = 0
for f_idx, f_p, e_idx, e_p in pairs:
    if len(f_p.runs) == len(e_p.runs):
        same_runs += 1
    else:
        diff_runs += 1

print(f"Same runs count: {same_runs}")
print(f"Different runs count: {diff_runs}")

# Inspect some diff runs
print("\nSample different run counts:")
count = 0
for f_idx, f_p, e_idx, e_p in pairs:
    if len(f_p.runs) != len(e_p.runs):
        print(f"FR[{f_idx}] ({len(f_p.runs)} runs): {repr(f_p.text[:40])} <===> EN[{e_idx}] ({len(e_p.runs)} runs): {repr(e_p.text[:40])}")
        count += 1
        if count >= 10:
            break
