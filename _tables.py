import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
p = r"C:\Users\HP\Downloads\PARCOURS TS\STAGE\STAGE TS STIC 2\rapport-stage\mon-rapport\Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx"
d = Document(p)
print("tables:", len(d.tables))
for i,t in enumerate(d.tables):
    print("--- table", i, "rows", len(t.rows), "cols", len(t.columns))
    print("   first row:", [c.text.strip()[:40] for c in t.rows[0].cells])
body = d.element.body
from docx.oxml.ns import qn
kids = [c.tag.split('}')[1] for c in body]
print()
print("body children count:", len(kids))
print("last 40 tags:", kids[-40:])
