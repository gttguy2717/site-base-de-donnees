import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc_en = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_ENGLISH.docx')
doc_fr = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

print(f"FR Paragraphs: {len(doc_fr.paragraphs)}")
print(f"EN Paragraphs: {len(doc_en.paragraphs)}")

# Let's inspect the sections in FR where Genius Pay appears:
# 1. P[71]: Figure 17 : Logo de la plateforme transactionnelle Genius Pay
# 2. P[393]: Ce diagramme de sequence presente le deroulement d'une transaction financiere electronique realisee via la passerelle Genius Pay...
# 3. P[558]: 3. Le Module de paiement en ligne securise (Genius Pay)
# 4. P[559]: Pour securiser les reglements...
# 5. P[560]: Infrastructure et passerelle...
# 6. P[562]: Validation et notifications...

# Let's see what is at P[380..400] and P[550..570] in FR and EN
print("\n--- Context of P[393] in FR ---")
for i in range(388, 400):
    print(f"FR[{i}]: {repr(doc_fr.paragraphs[i].text[:70])}")

print("\n--- Context around sequence diagrams in EN ---")
for i, p in enumerate(doc_en.paragraphs):
    if "sequence diagram" in p.text.lower():
        print(f"EN[{i}]: {repr(p.text[:70])}")

print("\n--- Context of P[558..565] in FR ---")
for i in range(555, 566):
    print(f"FR[{i}]: {repr(doc_fr.paragraphs[i].text[:70])}")

print("\n--- Context around Chapter 5 implementation modules in EN ---")
for i, p in enumerate(doc_en.paragraphs):
    if "module" in p.text.lower() or "payment" in p.text.lower() or "genius" in p.text.lower():
        print(f"EN[{i}]: {repr(p.text[:70])}")
