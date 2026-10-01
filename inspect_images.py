# -*- coding: utf-8 -*-
import sys; sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

# Check all images in FR doc
doc_fr = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')
print('=== FR doc image rId mapping ===')
for rId, rel in doc_fr.part.rels.items():
    if 'image' in rel.reltype.lower():
        print(f'  {rId} -> {rel.target_ref}')

print()
# Find Genius Pay logo context (paragraphs 536-545)
print('=== Context around paragraph 540 ===')
for i in range(536, 545):
    p = doc_fr.paragraphs[i]
    blips = p._element.findall('.//' + qn('a:blip'))
    has_img = len(blips) > 0
    rid_info = ''
    if has_img:
        rid_info = '[IMG rId=' + blips[0].get(qn('r:embed')) + '] '
    print(f'[{i}][{p.style.name}] {rid_info}{p.text[:100]}')

# Check English doc image rIds
doc_en = Document('Rapport_de_Stage_ENGLISH.docx')
print()
print('=== EN doc image rId mapping ===')
for rId, rel in doc_en.part.rels.items():
    if 'image' in rel.reltype.lower():
        print(f'  {rId} -> {rel.target_ref}')
