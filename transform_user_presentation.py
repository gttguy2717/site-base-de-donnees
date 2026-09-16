import os
import sys
from PIL import Image
import pptx
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

sys.stdout.reconfigure(encoding='utf-8')

# --- CHARTE VERT DU LOGO SOUTARAH GROUP (ZÉRO ORANGE) ---
LOGO_DARK_GREEN  = RGBColor(8, 44, 27)     # #082C1B
LOGO_BRAND_GREEN = RGBColor(40, 129, 52)   # #288134
LOGO_LEAF_LIME   = RGBColor(120, 183, 43)  # #78B72B
LIME_LIGHT_BG    = RGBColor(236, 248, 226) # #ECF8E2

def is_orange(rgb):
    """Détecte les nuances orangées/rousses."""
    r, g, b = rgb[0], rgb[1], rgb[2]
    return (r > 170 and g < 140 and b < 60) or (r > 210 and g < 160 and b < 80)

def replace_colors_in_shape(shape):
    """Remplace récursivement l'orange par le vert du logo Soutarah."""
    # 1. Fill
    if hasattr(shape, 'fill') and shape.fill.type == 1:
        try:
            rgb = shape.fill.fore_color.rgb
            if is_orange(rgb):
                shape.fill.solid()
                shape.fill.fore_color.rgb = LOGO_BRAND_GREEN
        except:
            pass

    # 2. Line
    if hasattr(shape, 'line') and shape.line.fill.type == 1:
        try:
            rgb = shape.line.color.rgb
            if is_orange(rgb):
                shape.line.color.rgb = LOGO_BRAND_GREEN
        except:
            pass

    # 3. Text
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                try:
                    rgb = r.font.color.rgb
                    if is_orange(rgb):
                        r.font.color.rgb = LOGO_LEAF_LIME
                except:
                    pass

    # 4. Group shapes
    if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            replace_colors_in_shape(child)

def place_image_in_bounds(slide, img_path, left, top, width, height):
    """Place une image avec centrage et conservation exacte du ratio d'aspect."""
    if not os.path.exists(img_path):
        print(f"Warning: {img_path} not found")
        return None
    with Image.open(img_path) as im:
        img_w, img_h = im.size
    
    img_aspect = img_w / img_h
    box_aspect = width / height

    if img_aspect > box_aspect:
        final_w = width
        final_h = width / img_aspect
        final_left = left
        final_top = top + (height - final_h) / 2
    else:
        final_h = height
        final_w = height * img_aspect
        final_top = top
        final_left = left + (width - final_w) / 2

    return slide.shapes.add_picture(img_path, final_left, final_top, width=final_w, height=final_h)

print("Helper functions defined successfully.")
