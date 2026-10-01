import sys
from docx import Document

doc_en = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_ENGLISH.docx')
doc_fr = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

print(f"FR doc paragraphs: {len(doc_fr.paragraphs)}, tables: {len(doc_fr.tables)}")
print(f"EN doc paragraphs: {len(doc_en.paragraphs)}, tables: {len(doc_en.tables)}")

# Check Sommaire in EN doc
print("\n=== SOMMAIRE EN DOC ===")
for i in range(35, 60):
    if i < len(doc_en.paragraphs):
        t = doc_en.paragraphs[i].text.strip()
        if t:
            print(f"[{i}] {t}")

# Check where Brevo / Hostinger are in EN doc
print("\n=== SEARCHING FOR BREVO / HOSTINGER IN EN DOC ===")
for i, p in enumerate(doc_en.paragraphs):
    t = p.text.strip().lower()
    if 'brevo' in t or 'hostinger' in t or 'genius' in t or 'payment' in t:
        print(f"[{i}] {p.text.strip()[:80]}")
