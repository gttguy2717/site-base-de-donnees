import docx
import json
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc_fr = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')
doc_en = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_ENGLISH.docx')

with open('fr_clean_474.json', encoding='utf-8') as f:
    fr_paras = json.load(f)

en_non_empty = [(i, p.text.strip()) for i, p in enumerate(doc_en.paragraphs) if p.text.strip()]

print(f"FR items: {len(fr_paras)}, EN items: {len(en_non_empty)}")

# Translations for the 6 Genius Pay paragraphs
genius_pay_trans = {
    71: "Figure 17 : Logo of the Genius Pay Transactional Platform\t - 40 -",
    393: "This sequence diagram details an electronic financial transaction processed via the Genius Pay gateway. The sequence begins when the user submits their payment details. The platform then securely transmits the transaction payload to the Genius Pay REST API, which handles validation with banking networks and Mobile Money operators (Wave, Orange Money, MTN MoMo). Upon receipt of the transaction outcome, Genius Pay dispatches an instantaneous, signed webhook notification back to our server, enabling automated, real-time booking confirmation.",
    558: "3. Secure Online Payment Module (Genius Pay)",
    559: "To secure transactions and automate order fulfillment, an online payment gateway was seamlessly integrated into the platform. This modern infrastructure replaces manual cash handling and wire transfers with instant digital settlements.",
    560: "Infrastructure and Gateway : Our Node.js backend initializes the transaction via the Genius Pay REST API, establishing a secured checkout session protected by mutual TLS and API keys. The client is seamlessly routed to the payment portal without exposing sensitive payment credentials to our application servers.",
    562: "Payment Validation and Webhook Notifications : Upon completion of bank authentication (3D-Secure or Mobile Money OTP), Genius Pay instantly triggers an asynchronous HTTPS Webhook notification to our Node.js server. The webhook payload contains the cryptographic transaction signature, transaction ID, paid amount, and payment status, allowing automatic reservation confirmation and invoice generation.",
}

exact_map = {}
fr_idx = 0
en_idx = 0

while fr_idx < len(fr_paras):
    f_item = fr_paras[fr_idx]
    f_p_idx = f_item['idx']
    f_txt = f_item['text'].strip()
    
    if f_p_idx in genius_pay_trans:
        exact_map[f_p_idx] = genius_pay_trans[f_p_idx]
        fr_idx += 1
        continue
        
    if en_idx < len(en_non_empty):
        e_orig_idx, e_txt = en_non_empty[en_idx]
        exact_map[f_p_idx] = e_txt
        fr_idx += 1
        en_idx += 1
    else:
        print(f"ERROR: Ran out of EN paragraphs at FR[{f_p_idx}]: {f_txt[:40]}")
        break

print(f"Total mapped: {len(exact_map)} / {len(fr_paras)}")

# Save to file
with open('exact_translations_474.json', 'w', encoding='utf-8') as f:
    json.dump({str(k): v for k, v in exact_map.items()}, f, ensure_ascii=False, indent=2)

print("Saved exact_translations_474.json successfully!")

# Verify alignment on critical sections
print("\n--- Spot Check Alignment ---")
for idx in [1, 2, 5, 6, 7, 29, 30, 31, 38, 55, 71, 87, 96, 97, 103, 104, 105, 106, 113, 114, 141, 142, 143, 144, 393, 558, 559, 560, 562, 658, 715, 716, 807]:
    if idx in exact_map:
        f_txt = doc_fr.paragraphs[idx].text.strip().replace('\n', ' ')
        e_txt = exact_map[idx].replace('\n', ' ')
        print(f"P[{idx:3d}]")
        print(f"  FR: {f_txt[:60]}")
        print(f"  EN: {e_txt[:60]}")
