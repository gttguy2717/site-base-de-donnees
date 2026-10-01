import docx

doc_en = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_ENGLISH.docx')
doc_fr = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

print("EN doc:")
print("  paragraphs:", len(doc_en.paragraphs))
print("  drawings:", len(doc_en._element.xpath('.//*[local-name()="drawing"]')))
print("  picts:", len(doc_en._element.xpath('.//*[local-name()="pict"]')))
print("  textboxes:", len(doc_en._element.xpath('.//*[local-name()="txbxContent"]')))

print("\nFR doc:")
print("  paragraphs:", len(doc_fr.paragraphs))
print("  drawings:", len(doc_fr._element.xpath('.//*[local-name()="drawing"]')))
print("  picts:", len(doc_fr._element.xpath('.//*[local-name()="pict"]')))
print("  textboxes:", len(doc_fr._element.xpath('.//*[local-name()="txbxContent"]')))

# Check chapter banners in EN doc textboxes
txbx_en = doc_en._element.xpath('.//*[local-name()="txbxContent"]')
print(f"\nEN textboxes: {len(txbx_en)}")
for i, tb in enumerate(txbx_en):
    t_nodes = tb.xpath('.//*[local-name()="t"]')
    txt = ' '.join([t.text for t in t_nodes if t.text]).strip()
    if 'chapitre' in txt.lower() or 'chapter' in txt.lower() or 'tableau' in txt.lower() or 'table' in txt.lower() or 'figure' in txt.lower():
        print(f"  [{i}]: {txt[:70]}")
