import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from build_common import (
    set_slide_background, add_header, add_footer, add_card, add_badge,
    create_section_divider, place_image_in_box, create_table,
    LOGO_DARK_GREEN, LOGO_BRAND_GREEN, LOGO_LEAF_LIME, LIME_LIGHT_BG, MINT_LIGHT_BG,
    DARK_CARD_BG, LIGHT_BG, CARD_BG, CARD_BORDER, TEXT_DARK, TEXT_BODY, TEXT_MUTED,
    TEXT_WHITE, FONT_HEADING, FONT_BODY
)

def build_slides_31_to_45(prs):
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 31 : SECTION 5 - DIAGRAMME DE SÉQUENCE 2 : RÉSERVATION VÉHICULE (FIGURE 7)
    # =========================================================================
    s31 = prs.slides.add_slide(blank_layout)
    set_slide_background(s31, prs, LIGHT_BG)
    add_header(s31, "5. CONCEPTION DU SYSTÈME • SÉQUENCE RÉSERVATION",
               "Réservation Mobile & Synchronisation Backend REST (Figure 7)",
               "Contrôle d'agenda en temps réel, émission du devis et notification Brevo")

    img_seq_res = "extracted_docx_images/image7.png"
    add_card(s31, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s31, img_seq_res, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    add_card(s31, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s31, Inches(9.25), Inches(2.1), Inches(2.8), Inches(0.32), "ANALYSE DE SÉQUENCE")

    tb_sr = s31.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_sr = tb_sr.text_frame
    tf_sr.word_wrap = True
    psr = tf_sr.paragraphs[0]
    psr.text = "Cinématique Réservation"
    psr.font.name = FONT_HEADING
    psr.font.size = Pt(13)
    psr.font.bold = True
    psr.font.color.rgb = TEXT_DARK

    psrb = tf_sr.add_paragraph()
    psrb.text = "1. L'application mobile soumet les dates choisies au point d'accès GET /api/vehicles/:id/availability.\n\n" \
                "2. Le serveur interroge la table Reservations pour détecter tout chevauchement de dates.\n\n" \
                "3. Après confirmation du client, l'API insère la réservation (statut PENDING) et génère le devis lié.\n\n" \
                "4. Envoi automatisé d'un email récapitulatif au client et d'une alerte à l'administrateur via Brevo."
    psrb.font.name = FONT_BODY
    psrb.font.size = Pt(10)
    psrb.font.color.rgb = TEXT_BODY
    psrb.space_before = Pt(8)

    add_footer(s31, 31)

    # =========================================================================
    # SLIDE 32 : SECTION 5 - DIAGRAMME DE SÉQUENCE 3 : PAIEMENT EN LIGNE (FIGURE 6)
    # =========================================================================
    s32 = prs.slides.add_slide(blank_layout)
    set_slide_background(s32, prs, LIGHT_BG)
    add_header(s32, "5. CONCEPTION DU SYSTÈME • SÉQUENCE PAIEMENT",
               "Paiement en Ligne Sécurisé via Passerelle Genius Pay (Figure 6)",
               "Intégration du guichet de paiement électronique multi-opérateurs (Cartes et Mobile Money)")

    img_seq_pay = "extracted_docx_images/image6.png"
    add_card(s32, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s32, img_seq_pay, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    add_card(s32, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s32, Inches(9.25), Inches(2.1), Inches(2.8), Inches(0.32), "ANALYSE DE SÉQUENCE")

    tb_sp = s32.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_sp = tb_sp.text_frame
    tf_sp.word_wrap = True
    psp = tf_sp.paragraphs[0]
    psp.text = "Flux Monétique Sécurisé"
    psp.font.name = FONT_HEADING
    psp.font.size = Pt(13)
    psp.font.bold = True
    psp.font.color.rgb = TEXT_DARK

    pspb = tf_sp.add_paragraph()
    pspb.text = "1. À la validation du panier, le client initie la transaction financière.\n\n" \
                "2. Le serveur API contacte la passerelle Genius Pay pour créer une session chiffrée et générer l'URL de paiement.\n\n" \
                "3. Le client règle sur l'interface sécurisée (Orange Money, MTN MoMo, Wave ou Carte bancaire).\n\n" \
                "4. Un webhook sécurisé notifie l'API pour passer la commande au statut 'PAID' et Brevo expédie le reçu officiel."
    pspb.font.name = FONT_BODY
    pspb.font.size = Pt(10)
    pspb.font.color.rgb = TEXT_BODY
    pspb.space_before = Pt(8)

    add_footer(s32, 32)

    # =========================================================================
    # SLIDE 33 : SECTION 5 - DIAGRAMME DE SÉQUENCE 4 : ACHAT PRODUIT (FIGURE 8)
    # =========================================================================
    s33 = prs.slides.add_slide(blank_layout)
    set_slide_background(s33, prs, LIGHT_BG)
    add_header(s33, "5. CONCEPTION DU SYSTÈME • SÉQUENCE ACHAT",
               "Processus d'Achat Produit & Traitement Commercial (Figure 8)",
               "Enregistrement de la commande de négoce, validation du devis et ordonnancement logistique")

    img_seq_ach = "extracted_docx_images/image8.png"
    add_card(s33, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s33, img_seq_ach, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    add_card(s33, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s33, Inches(9.25), Inches(2.1), Inches(2.8), Inches(0.32), "ANALYSE DE SÉQUENCE")

    tb_sa = s33.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_sa = tb_sa.text_frame
    tf_sa.word_wrap = True
    psa = tf_sa.paragraphs[0]
    psa.text = "Flux Commercial Négoce"
    psa.font.name = FONT_HEADING
    psa.font.size = Pt(13)
    psa.font.bold = True
    psa.font.color.rgb = TEXT_DARK

    psab = tf_sa.add_paragraph()
    psab.text = "1. Le client valide les articles de négoce sélectionnés dans son panier.\n\n" \
                "2. L'API REST insère la commande en base de données avec l'ensemble des lignes de détails chiffrées.\n\n" \
                "3. Le devis est validé automatiquement et Brevo transmet la confirmation avec le bon de commande en pièce jointe.\n\n" \
                "4. L'administrateur consulte la notification sur son tableau de bord et marque le devis comme lu pour préparer le colis."
    psab.font.name = FONT_BODY
    psab.font.size = Pt(10)
    psab.font.color.rgb = TEXT_BODY
    psab.space_before = Pt(8)

    add_footer(s33, 33)

    # =========================================================================
    # SLIDE 34 : SECTION 5 - DIAGRAMME DE CLASSES (FIGURE 10)
    # =========================================================================
    s34 = prs.slides.add_slide(blank_layout)
    set_slide_background(s34, prs, LIGHT_BG)
    add_header(s34, "5. CONCEPTION DU SYSTÈME • MODÉLISATION STATIQUE",
               "Diagramme de Classes du Système d'Information (Figure 10)",
               "Structure statique des entités métier et correspondances directes avec les modèles ORM")

    img_classes = "extracted_docx_images/image10.png"
    add_card(s34, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s34, img_classes, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    add_card(s34, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s34, Inches(9.25), Inches(2.1), Inches(2.8), Inches(0.32), "ANALYSE DE STRUCTURE")

    tb_cl = s34.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_cl = tb_cl.text_frame
    tf_cl.word_wrap = True
    pcl = tf_cl.paragraphs[0]
    pcl.text = "Entités & Cardinalités"
    pcl.font.name = FONT_HEADING
    pcl.font.size = Pt(13)
    pcl.font.bold = True
    pcl.font.color.rgb = TEXT_DARK

    pclb = tf_cl.add_paragraph()
    pclb.text = "• User & Client : relation 1:1 pour dissocier compte système et informations fiscales (Entreprise/Particulier).\n\n" \
                "• Vehicle & Reservation : relation 1:N avec vérification d'intégrité sur les intervalles temporels.\n\n" \
                "• Quote (Devis) : entité pivot liée au Client, agrégeant les QuoteItems (articles négoce et locations).\n\n" \
                "• Service & Product : catalogue modulaire avec tarification dynamique selon profil."
    pclb.font.name = FONT_BODY
    pclb.font.size = Pt(10)
    pclb.font.color.rgb = TEXT_BODY
    pclb.space_before = Pt(8)

    add_footer(s34, 34)

    # =========================================================================
    # SLIDE 35 : SECTION 5 - ARCHITECTURE GÉNÉRALE DU SYSTÈME (FIGURE 18)
    # =========================================================================
    s35 = prs.slides.add_slide(blank_layout)
    set_slide_background(s35, prs, LIGHT_BG)
    add_header(s35, "5. CONCEPTION DU SYSTÈME • ARCHITECTURE LOGICIELLE",
               "Schéma de l'Architecture Générale du Système (Figure 18)",
               "Architecture distribuée en 3-tiers découplée : Clients Web/Mobile, API REST et Données")

    img_arch = "extracted_docx_images/image18.png"
    add_card(s35, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s35, img_arch, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    add_card(s35, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s35, Inches(9.25), Inches(2.1), Inches(2.8), Inches(0.32), "DÉCOUPAGE EN 3 TIERS")

    tb_ar = s35.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_ar = tb_ar.text_frame
    tf_ar.word_wrap = True
    par = tf_ar.paragraphs[0]
    par.text = "Organisation des Couches"
    par.font.name = FONT_HEADING
    par.font.size = Pt(13)
    par.font.bold = True
    par.font.color.rgb = TEXT_DARK

    parb = tf_ar.add_paragraph()
    parb.text = "• Niveau Présentation : React.js (Web 6 Pôles) & React Native Expo (Mobile 21 Véhicules).\n\n" \
                "• Niveau Réseau & Transport : communications sécurisées HTTPS et échanges RESTful en JSON pur.\n\n" \
                "• Niveau Métier (Serveur) : API Node.js / Express assurant l'authentification, les règles anti-collision et le moteur PDF.\n\n" \
                "• Niveau Données : SGBD MySQL piloté par Sequelize ORM avec transactions ACID."
    parb.font.name = FONT_BODY
    parb.font.size = Pt(10)
    parb.font.color.rgb = TEXT_BODY
    parb.space_before = Pt(8)

    add_footer(s35, 35)

    # =========================================================================
    # SLIDE 36 : SECTION 5 - TECHNOLOGIES UTILISÉES 1 : FRONTEND WEB & MOBILE
    # =========================================================================
    s36 = prs.slides.add_slide(blank_layout)
    set_slide_background(s36, prs, LIGHT_BG)
    add_header(s36, "5. CONCEPTION DU SYSTÈME • CHOIX TECHNOLOGIQUES (1/3)",
               "Technologies Frontend : Plateforme Web & Application Mobile",
               "Frameworks modernes garantissant performance, réactivité et modularité de code")

    # Carte Web
    add_card(s36, Inches(0.8), Inches(1.85), Inches(5.7), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s36, Inches(1.05), Inches(2.1), Inches(3.0), Inches(0.32), "FRONTEND WEB : REACT.JS & TAILWIND")

    img_react = "extracted_docx_images/image13.jpg"
    if os.path.exists(img_react):
        place_image_in_box(s36, img_react, Inches(1.05), Inches(2.55), Inches(1.2), Inches(1.0))

    tb_w = s36.shapes.add_textbox(Inches(2.4), Inches(2.5), Inches(3.8), Inches(1.0))
    tf_w = tb_w.text_frame
    pw = tf_w.paragraphs[0]
    pw.text = "React.js v18 & Tailwind CSS"
    pw.font.name = FONT_HEADING
    pw.font.size = Pt(13.5)
    pw.font.bold = True
    pw.font.color.rgb = TEXT_DARK

    tb_wb = s36.shapes.add_textbox(Inches(1.05), Inches(3.65), Inches(5.2), Inches(2.9))
    tf_wb = tb_wb.text_frame
    tf_wb.word_wrap = True
    pwb = tf_wb.paragraphs[0]
    pwb.text = "• Architecture par Composants : modularité extrême du catalogue, du panier interactif et de la vitrine.\n\n" \
               "• Virtual DOM : fluidité exceptionnelle lors des filtres multicritères et mises à jour en direct.\n\n" \
               "• Tailwind CSS : design system sur mesure aligné à 100% sur la charte graphique verte de SOUTARAH GROUP.\n\n" \
               "• Vite Build Tool : temps de compilation instantané et bundle de production optimisé."
    pwb.font.name = FONT_BODY
    pwb.font.size = Pt(10.5)
    pwb.font.color.rgb = TEXT_BODY

    # Carte Mobile
    add_card(s36, Inches(6.8), Inches(1.85), Inches(5.7), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s36, Inches(7.05), Inches(2.1), Inches(3.2), Inches(0.32), "MOBILE : REACT NATIVE & EXPO")

    img_rn = "extracted_docx_images/image14.png"
    if os.path.exists(img_rn):
        place_image_in_box(s36, img_rn, Inches(7.05), Inches(2.55), Inches(1.5), Inches(1.0))

    tb_m = s36.shapes.add_textbox(Inches(8.7), Inches(2.5), Inches(3.5), Inches(1.0))
    tf_m = tb_m.text_frame
    pm = tf_m.paragraphs[0]
    pm.text = "React Native & Expo"
    pm.font.name = FONT_HEADING
    pm.font.size = Pt(13.5)
    pm.font.bold = True
    pm.font.color.rgb = TEXT_DARK

    tb_mb = s36.shapes.add_textbox(Inches(7.05), Inches(3.65), Inches(5.2), Inches(2.9))
    tf_mb = tb_mb.text_frame
    tf_mb.word_wrap = True
    pmb = tf_mb.paragraphs[0]
    pmb.text = "• Codebase Unique : un seul code source compilé nativement pour les plateformes Android et iOS.\n\n" \
               "• Écosystème Expo : accélération majeure du cycle de test et gestion fluide des permissions matérielles.\n\n" \
               "• Ergonomie Mobile Native : navigation fluide par onglets, sélecteurs tactiles de dates et fluidité 60 FPS.\n\n" \
               "• Visionneuse PDF Intégrée : ouverture et partage direct des devis officiels sur smartphone."
    pmb.font.name = FONT_BODY
    pmb.font.size = Pt(10.5)
    pmb.font.color.rgb = TEXT_BODY

    add_footer(s36, 36)

    # =========================================================================
    # SLIDE 37 : SECTION 5 - TECHNOLOGIES UTILISÉES 2 : BACKEND & SGBD
    # =========================================================================
    s37 = prs.slides.add_slide(blank_layout)
    set_slide_background(s37, prs, LIGHT_BG)
    add_header(s37, "5. CONCEPTION DU SYSTÈME • CHOIX TECHNOLOGIQUES (2/3)",
               "Technologies Serveur : Node.js, Express.js & MySQL",
               "Architecture Full-Stack JavaScript robuste et moteur relationnel éprouvé")

    # Carte Backend
    add_card(s37, Inches(0.8), Inches(1.85), Inches(5.7), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s37, Inches(1.05), Inches(2.1), Inches(3.0), Inches(0.32), "SERVEUR : NODE.JS & EXPRESS.JS")

    img_node = "extracted_docx_images/image15.png"
    if os.path.exists(img_node):
        place_image_in_box(s37, img_node, Inches(1.05), Inches(2.55), Inches(1.3), Inches(1.0))

    tb_n = s37.shapes.add_textbox(Inches(2.5), Inches(2.5), Inches(3.8), Inches(1.0))
    tf_n = tb_n.text_frame
    pn = tf_n.paragraphs[0]
    pn.text = "Node.js & Express.js"
    pn.font.name = FONT_HEADING
    pn.font.size = Pt(13.5)
    pn.font.bold = True
    pn.font.color.rgb = TEXT_DARK

    tb_nb = s37.shapes.add_textbox(Inches(1.05), Inches(3.65), Inches(5.2), Inches(2.9))
    tf_nb = tb_nb.text_frame
    tf_nb.word_wrap = True
    pnb = tf_nb.paragraphs[0]
    pnb.text = "• Moteur Asynchrone : gestion non bloquante par boucle d'événements, idéale pour les requêtes simultanées.\n\n" \
               "• Homogénéité Full-Stack JS : partage des types de données et compétences JavaScript unifiées du front au back.\n\n" \
               "• Sécurité & Contrôle : middlewares dédiés pour la validation JWT, le hachage Bcrypt et la protection CORS.\n\n" \
               "• Générateur PDFKit : moteur serveur compilant les devis PDF 2 pages avec en-tête juridique et logos."
    pnb.font.name = FONT_BODY
    pnb.font.size = Pt(10.5)
    pnb.font.color.rgb = TEXT_BODY

    # Carte SGBD
    add_card(s37, Inches(6.8), Inches(1.85), Inches(5.7), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s37, Inches(7.05), Inches(2.1), Inches(3.2), Inches(0.32), "PERSISTANCE : MYSQL & SEQUELIZE")

    img_db_schema = "extracted_docx_images/image19.png"
    if os.path.exists(img_db_schema):
        place_image_in_box(s37, img_db_schema, Inches(7.05), Inches(2.55), Inches(1.8), Inches(1.0))

    tb_my = s37.shapes.add_textbox(Inches(9.0), Inches(2.5), Inches(3.2), Inches(1.0))
    tf_my = tb_my.text_frame
    pmy = tf_my.paragraphs[0]
    pmy.text = "MySQL & Sequelize ORM"
    pmy.font.name = FONT_HEADING
    pmy.font.size = Pt(13.5)
    pmy.font.bold = True
    pmy.font.color.rgb = TEXT_DARK

    tb_myb = s37.shapes.add_textbox(Inches(7.05), Inches(3.65), Inches(5.2), Inches(2.9))
    tf_myb = tb_myb.text_frame
    tf_myb.word_wrap = True
    pmyb = tf_myb.paragraphs[0]
    pmyb.text = "• Rigueur Relationnelle : respect strict des clés primaires, étrangères et contraintes d'intégrité.\n\n" \
                "• Transactions ACID : sécurisation absolue des écritures concurrentes lors des réservations de véhicules.\n\n" \
                "• ORM Sequelize : abstraction orientée objet facilitant les migrations, requêtes complexes et relations.\n\n" \
                "• Parfaite Adéquation Hébergement : compatibilité 100% native avec l'infrastructure Hostinger acquise."
    pmyb.font.name = FONT_BODY
    pmyb.font.size = Pt(10.5)
    pmyb.font.color.rgb = TEXT_BODY

    add_footer(s37, 37)

    # =========================================================================
    # SLIDE 38 : SECTION 5 - TECHNOLOGIES UTILISÉES 3 : OUTILS & SERVICES TIERS
    # =========================================================================
    s38 = prs.slides.add_slide(blank_layout)
    set_slide_background(s38, prs, LIGHT_BG)
    add_header(s38, "5. CONCEPTION DU SYSTÈME • CHOIX TECHNOLOGIQUES (3/3)",
               "Outils de Conception, Services Tiers & Déploiement",
               "Une boîte à outils moderne au service d'une mise en production industrielle")

    outils = [
        ("VS CODE & FIGMA", "IDE & Maquettage UX/UI",
         "• VS Code : environnement de développement léger avec plugins TypeScript, ESLint et Git intégrés.\n• Figma : conception collaborative des wireframes et prototypes graphiques avant codage."),
        ("BREVO (EX-SENDINBLUE)", "Notifications Transactionnelles",
         "• Passerelle SMTP / API Cloud garantissant une délivrabilité optimale sans passage en spam.\n• Automatisation des e-mails : confirmations de réservation, devis PDF attachés et alertes admin."),
        ("GENIUS PAY GATEWAY", "Passerelle de Paiement Sécurisée",
         "• Agrégateur monétique ivoirien supportant Wave, Orange Money, MTN MoMo et cartes Visa/Mastercard.\n• Sécurisation 3D Secure et notification synchrone par webhooks chiffrés."),
        ("HOSTINGER CLOUD", "Hébergement Web & Base de Données",
         "• Serveur VPS / Cloud haute disponibilité garantissant un temps de disponibilité (uptime) de 99.9%.\n• Certificats de sécurité SSL/TLS Let's Encrypt et sauvegardes automatiques journalières.")
    ]

    for idx, (badge_t, title, desc) in enumerate(outils):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8) if col == 0 else Inches(6.8)
        top = Inches(1.85) + row * Inches(2.45)
        width = Inches(5.7)
        height = Inches(2.25)

        add_card(s38, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
        add_badge(s38, left + Inches(0.25), top + Inches(0.25), Inches(3.2), Inches(0.32), badge_t)

        tb = s38.shapes.add_textbox(left + Inches(0.25), top + Inches(0.65), Inches(5.2), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_BODY
        pd.space_before = Pt(4)

    add_footer(s38, 38)

    # =========================================================================
    # SLIDE 39 : SECTION 5 - ÉTUDE COMPARATIVE DES SGBD (TABLEAU 9)
    # =========================================================================
    s39 = prs.slides.add_slide(blank_layout)
    set_slide_background(s39, prs, LIGHT_BG)
    add_header(s39, "5. CONCEPTION DU SYSTÈME • SÉLECTION DU SGBD",
               "Comparatif des Principaux SGBD & Choix de MySQL (Tableau 9)",
               "Analyse comparative rigoureuse des moteurs de bases de données étudiés")

    headers_sgbd = ["Critères", "MySQL (Retenu)", "PostgreSQL", "SQL Server", "Oracle", "MongoDB"]
    rows_sgbd = [
        ["Type de système", "Relationnel, Open Source", "Relationnel, Open Source", "Relationnel, Propriétaire", "Relationnel, Propriétaire", "NoSQL, Documents"],
        ["Modèle de données", "Relationnel structuré", "Relationnel structuré", "Relationnel structuré", "Relationnel structuré", "Collections JSON flexibles"],
        ["Performance", "Très bonne sur lectures", "Très bonne requêtes avancées", "Très bonne", "Excellente grands comptes", "Excellente volumes bruts"],
        ["Coût de licence", "Gratuit (GPL)", "Gratuit (BSD)", "Payant (très onéreux)", "Payant (très onéreux)", "Gratuit / Cloud payant"],
        ["Facilité d'usage", "Élevée (communauté vaste)", "Moyenne à élevée", "Élevée (écosystème MS)", "Complexe (administration)", "Élevée"],
        ["Compatibilité hébergement", "OUI (Supporté à 100%)", "Non disponible (offre souscrite)", "Non supporté", "Non supporté", "Non supporté"]
    ]
    col_w_sgbd = [Inches(2.2), Inches(2.2), Inches(2.1), Inches(1.8), Inches(1.8), Inches(1.63)]
    create_table(s39, Inches(0.8), Inches(1.85), Inches(11.73), Inches(4.8), headers_sgbd, rows_sgbd, col_w_sgbd)

    add_footer(s39, 39)

    # =========================================================================
    # SLIDE 40 : INTERCALAIRE SECTION 6 : RÉALISATION
    # =========================================================================
    create_section_divider(prs, 40, 6, "RÉALISATION",
                           "Mise en œuvre concrète de l'assistant IA, intégration du service de notifications Brevo, moteur de devis PDF et démonstration commentée des interfaces.")

    # =========================================================================
    # SLIDE 41 : SECTION 6 - ASSISTANT IA INTELLIGENT
    # =========================================================================
    s41 = prs.slides.add_slide(blank_layout)
    set_slide_background(s41, prs, LIGHT_BG)
    add_header(s41, "6. RÉALISATION • ASSISTANT INTELLIGENT",
               "Implémentation de l’Assistant Virtuel IA",
               "Un module conversationnel interactif au service de l'orientation client et du support 24/7")

    cards_ia = [
        ("ORIENTATION 24/7", "Aiguillage Immédiat des Visiteurs",
         "• Présent directement sur le portail Web sous forme de widget conversationnel dynamique.\n• Répond instantanément aux questions des usagers sur les 6 pôles d'activités.\n• Oriente avec précision vers le bon pôle : réservation de véhicule, négoce ou devis sur mesure.\n• Réduit drastiquement l'hésitation des primo-visiteurs et augmente la conversion."),
        ("CONNAISSANCE MÉTIER", "Prompt Spécialisé SOUTARAH",
         "• Connaissance exhaustive des valeurs S.A.P.E et de la flotte des 21 véhicules disponibles.\n• Explication transparente des conditions de location (tarifs Abidjan / Hors Abidjan, chauffeur).\n• Clarification des procédures de devis pro-forma et pièces administratives requises.\n• Traitement des requêtes en langage naturel en français."),
        ("PILOTAGE COMMERCIAL", "Assistance à l'Administration",
         "• Capacité d'interfaçage avec le tableau de bord pour synthétiser les demandes récurrentes.\n• Identification des interrogations fréquentes pour affiner les fiches produits du catalogue.\n• Gain de temps substantiel pour le secrétariat commercial libéré des questions répétitives.\n• Évolution prête pour l'intégration de modèles LLM de dernière génération.")
    ]

    for i, (b_t, title, body) in enumerate(cards_ia):
        left = Inches(0.8) + i * Inches(3.95)
        top = Inches(1.85)
        width = Inches(3.8)
        height = Inches(4.9)
        add_card(s41, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
        add_badge(s41, left + Inches(0.25), top + Inches(0.25), Inches(3.2), Inches(0.32), b_t)

        tb = s41.shapes.add_textbox(left + Inches(0.25), top + Inches(0.75), Inches(3.3), Inches(3.9))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13.5)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK

        pb = tf.add_paragraph()
        pb.text = body
        pb.font.name = FONT_BODY
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = TEXT_BODY
        pb.space_before = Pt(10)

    add_footer(s41, 41)

    # =========================================================================
    # SLIDE 42 : SECTION 6 - NOTIFICATIONS TRANSACTIONNELLES BREVO
    # =========================================================================
    s42 = prs.slides.add_slide(blank_layout)
    set_slide_background(s42, prs, LIGHT_BG)
    add_header(s42, "6. RÉALISATION • MESSAGERIE TRANSACTIONNELLE",
               "Intégration du Service de Notifications Transactionnelles",
               "Chaîne de communication automatisée par courriel via l'API Cloud Brevo (SMTP)")

    cards_notif = [
        ("NOTIFICATION CLIENT", "Accusé de Réception Instantané",
         "• Déclenchement automatique dès la soumission d'une réservation ou validation de devis.\n• Envoi d'un courriel HTML soigné aux couleurs de la marque avec récapitulatif détaillé.\n• Pièce jointe intégrée : devis pro-forma officiel PDF 2 pages prêt pour signature.\n• Confirmation de l'enrôlement et e-mails de bienvenue lors de la création de compte."),
        ("ALERTE ADMINISTRATEUR", "Prise en Charge Immédiate",
         "• Alerte instantanée expédiée à la direction commerciale à chaque nouvelle commande.\n• Fiche synthétique : nom du client, véhicule ou produit sollicité, montant estimé HT/TTC.\n• Lien direct permettant d'ouvrir la fiche dans le back-office pour marquer le devis comme lu.\n• Suppression intégrale des retards de traitement des dossiers clients."),
        ("INFRASTRUCTURE BREVO", "Fiabilité & Délivrabilité",
         "• Intégration via API REST sécurisée avec clé API chiffrée dans les variables d'environnement.\n• Utilisation de templates transactionnels responsives adaptés aux messageries mobiles.\n• Suivi en temps réel de l'état de délivrance (ouvertures, clics, réceptions confirmées).\n• Conformité stricte aux standards anti-spam (SPF, DKIM, DMARC configurés).")
    ]

    for i, (b_t, title, body) in enumerate(cards_notif):
        left = Inches(0.8) + i * Inches(3.95)
        top = Inches(1.85)
        width = Inches(3.8)
        height = Inches(4.9)
        add_card(s42, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
        add_badge(s42, left + Inches(0.25), top + Inches(0.25), Inches(3.2), Inches(0.32), b_t)

        tb = s42.shapes.add_textbox(left + Inches(0.25), top + Inches(0.75), Inches(3.3), Inches(3.9))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13.5)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK

        pb = tf.add_paragraph()
        pb.text = body
        pb.font.name = FONT_BODY
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = TEXT_BODY
        pb.space_before = Pt(10)

    add_footer(s42, 42)

    # =========================================================================
    # SLIDE 43 : SECTION 6 - MOTEUR DE DEVIS PRO-FORMA PDF
    # =========================================================================
    s43 = prs.slides.add_slide(blank_layout)
    set_slide_background(s43, prs, LIGHT_BG)
    add_header(s43, "6. RÉALISATION • MOTEUR DE FACTURATION",
               "Génération Automatisée du Devis Pro-Forma PDF Officiel",
               "Moteur serveur PDFKit garantissant la stricte conformité fiscale et juridique ivoirienne")

    # Carte Gauche : Normes fiscales
    add_card(s43, Inches(0.8), Inches(1.85), Inches(5.7), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s43, Inches(1.05), Inches(2.1), Inches(3.0), Inches(0.32), "MOTEUR DE CALCUL FISCAL")

    tb_pdf1 = s43.shapes.add_textbox(Inches(1.05), Inches(2.55), Inches(5.2), Inches(4.0))
    tf_pdf1 = tb_pdf1.text_frame
    tf_pdf1.word_wrap = True
    p1 = tf_pdf1.paragraphs[0]
    p1.text = "Règles Fiscale & Légales Intégrées"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_DARK

    p1b = tf_pdf1.add_paragraph()
    p1b.text = "• Numérotation Normalisée : génération d'une référence unique horodatée (ex: DMD-2026-XXXX).\n\n" \
               "• Calcul Automatisé des Taxes Réglementaires :\n" \
               "  - Montant Total Hors Taxes (HT)\n" \
               "  - Taxe sur la Valeur Ajoutée (TVA) : 18%\n" \
               "  - Taxe de Développement Touristique (TDT) : 2.5% (applicable sur les locations de véhicules)\n" \
               "  - Total Toutes Taxes Comprises (TTC) en FCFA.\n\n" \
               "• Barème Kilométrique & Destination : intégration automatique du supplément 'Hors Abidjan' et du forfait chauffeur."
    p1b.font.name = FONT_BODY
    p1b.font.size = Pt(10.5)
    p1b.font.color.rgb = TEXT_BODY
    p1b.space_before = Pt(8)

    # Carte Droite : Structure du document PDF sur 2 pages
    add_card(s43, Inches(6.8), Inches(1.85), Inches(5.7), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s43, Inches(7.05), Inches(2.1), Inches(3.2), Inches(0.32), "DOCUMENT PDF 2 PAGES NORMALISÉ")

    tb_pdf2 = s43.shapes.add_textbox(Inches(7.05), Inches(2.55), Inches(5.2), Inches(4.0))
    tf_pdf2 = tb_pdf2.text_frame
    tf_pdf2.word_wrap = True
    p2 = tf_pdf2.paragraphs[0]
    p2.text = "Agencement Professionnel du Devis"
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_DARK

    p2b = tf_pdf2.add_paragraph()
    p2b.text = "• Page 1 : Devis Pro-Forma Contractuel\n" \
               "  - En-tête officiel SOUTARAH GROUP (Logo haute résolution, RCCM, CC, coordonnées bancaires BNI).\n" \
               "  - Bloc Client certifié (raison sociale, contact, adresse).\n" \
               "  - Tableau chiffré des prestations avec quantités, prix unitaires et totaux.\n" \
               "  - Zone de signature et cachet commercial.\n\n" \
               "• Page 2 : Fiches Techniques & Visuels HD\n" \
               "  - Visuel haute définition du ou des véhicules réservés.\n" \
               "  - Spécifications techniques détaillées (motorisation, transmission, options)."
    p2b.font.name = FONT_BODY
    p2b.font.size = Pt(10.5)
    p2b.font.color.rgb = TEXT_BODY
    p2b.space_before = Pt(8)

    add_footer(s43, 43)

    # =========================================================================
    # SLIDE 44 : SECTION 6 - INTERFACES WEB 1 : ACCUEIL VITRINE (FIGURE 20)
    # =========================================================================
    s44 = prs.slides.add_slide(blank_layout)
    set_slide_background(s44, prs, LIGHT_BG)
    add_header(s44, "6. RÉALISATION • INTERFACES WEB (1/4)",
               "Portail Web : Page d'Accueil & Présentation des Pôles (Figure 20)",
               "Une vitrine digitale immersive valorisant les activités multisectorielles de SOUTARAH")

    img_web_home = "extracted_docx_images/image20.png"
    add_card(s44, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s44, img_web_home, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    add_card(s44, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s44, Inches(9.25), Inches(2.1), Inches(2.8), Inches(0.32), "ANALYSE DE L'INTERFACE")

    tb_wh = s44.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_wh = tb_wh.text_frame
    tf_wh.word_wrap = True
    pwh = tf_wh.paragraphs[0]
    pwh.text = "Caractéristiques Clés"
    pwh.font.name = FONT_HEADING
    pwh.font.size = Pt(13)
    pwh.font.bold = True
    pwh.font.color.rgb = TEXT_DARK

    pwhb = tf_wh.add_paragraph()
    pwhb.text = "• Header dynamique avec navigation par pôle, accès direct au panier et bouton de connexion.\n\n" \
                "• Section Hero percutante avec proposition de valeur claire et appel à l'action immédiat.\n\n" \
                "• Présentation interactive des 6 pôles avec cartes visuelles et descriptifs des compétences.\n\n" \
                "• Intégration du widget de chat avec l'assistant virtuel IA en bas à droite de l'écran."
    pwhb.font.name = FONT_BODY
    pwhb.font.size = Pt(10)
    pwhb.font.color.rgb = TEXT_BODY
    pwhb.space_before = Pt(8)

    add_footer(s44, 44)

    # =========================================================================
    # SLIDE 45 : SECTION 6 - INTERFACES WEB 2 : CONNEXION & CATALOGUE (FIGURES 21 & 22)
    # =========================================================================
    s45 = prs.slides.add_slide(blank_layout)
    set_slide_background(s45, prs, LIGHT_BG)
    add_header(s45, "6. RÉALISATION • INTERFACES WEB (2/4)",
               "Authentification Client & Catalogue Produits (Figures 21 & 22)",
               "Parcours d'accès sécurisé et navigation ergonomique dans l'offre commerciale")

    # Image Connexion (Figure 21)
    img_auth = "extracted_docx_images/image21.png"
    add_card(s45, Inches(0.8), Inches(1.85), Inches(5.7), Inches(3.3), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s45, img_auth, Inches(0.9), Inches(1.95), Inches(5.5), Inches(3.1))

    # Image Catalogue (Figure 22)
    img_cat = "extracted_docx_images/image22.png"
    add_card(s45, Inches(6.8), Inches(1.85), Inches(5.7), Inches(3.3), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s45, img_cat, Inches(6.9), Inches(1.95), Inches(5.5), Inches(3.1))

    # 2 Cartes descriptives bas
    add_card(s45, Inches(0.8), Inches(5.3), Inches(5.7), Inches(1.5), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    tb_c1 = s45.shapes.add_textbox(Inches(1.0), Inches(5.4), Inches(5.3), Inches(1.3))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    pc1 = tf_c1.paragraphs[0]
    pc1.text = "Figure 21 : Authentification & Inscription"
    pc1.font.name = FONT_HEADING
    pc1.font.size = Pt(11.5)
    pc1.font.bold = True
    pc1.font.color.rgb = TEXT_DARK
    pc1b = tf_c1.add_paragraph()
    pc1b.text = "Formulaire sécurisé avec contrôle des formats, choix du profil (Particulier/Entreprise) et émission instantanée de session JWT."
    pc1b.font.name = FONT_BODY
    pc1b.font.size = Pt(9.5)
    pc1b.font.color.rgb = TEXT_BODY

    add_card(s45, Inches(6.8), Inches(5.3), Inches(5.7), Inches(1.5), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    tb_c2 = s45.shapes.add_textbox(Inches(7.0), Inches(5.4), Inches(5.3), Inches(1.3))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    pc2 = tf_c2.paragraphs[0]
    pc2.text = "Figure 22 : Catalogue Multi-Pôles"
    pc2.font.name = FONT_HEADING
    pc2.font.size = Pt(11.5)
    pc2.font.bold = True
    pc2.font.color.rgb = TEXT_DARK
    pc2b = tf_c2.add_paragraph()
    pc2b.text = "Filtres dynamiques par catégorie, visuels produits haute résolution, fiches techniques détaillées et bouton d'ajout au panier."
    pc2b.font.name = FONT_BODY
    pc2b.font.size = Pt(9.5)
    pc2b.font.color.rgb = TEXT_BODY

    add_footer(s45, 45)

    print("Slides 31 to 45 successfully built.")
