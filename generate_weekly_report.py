import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_weekly_report():
    doc = docx.Document()
    
    # Header
    p_logo = doc.add_paragraph()
    r_logo = p_logo.add_run("SOUTARAH GROUP")
    r_logo.bold = True
    r_logo.font.size = Pt(14)
    r_logo.font.color.rgb = RGBColor(21, 128, 61)
    
    doc.add_paragraph()
    
    # TITLE
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("RAPPORT DE SEMAINE")
    r_title.bold = True
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(31, 73, 125)
    
    doc.add_paragraph()
    
    # Meta info
    p_meta = doc.add_paragraph()
    p_meta.add_run("Période : ").bold = True
    p_meta.add_run("Du lundi 17/08/2026 au samedi 22/08/2026\n")
    p_meta.add_run("Lieu : ").bold = True
    p_meta.add_run("SOUTARAH GROUP\n")
    p_meta.add_run("Projet : ").bold = True
    p_meta.add_run("Développement du Backend (API REST) et de la Base de données de la plateforme.")
    
    doc.add_paragraph()
    
    # Activités réalisées
    p_act_title = doc.add_paragraph()
    r_act_title = p_act_title.add_run("Activités réalisées")
    r_act_title.bold = True
    r_act_title.font.size = Pt(12)
    r_act_title.font.color.rgb = RGBColor(31, 73, 125)
    
    doc.add_paragraph("Au cours de cette semaine, suite à la validation de l'architecture, j'ai achevé la phase de modélisation conceptuelle et j'ai entamé le développement du Backend (API REST) ainsi que la structuration de la base de données PostgreSQL de la plateforme.")
    
    # Lundi
    p = doc.add_paragraph()
    r = p.add_run("Lundi 17/08/2026")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 73, 125)
    doc.add_paragraph("- Finalisation de la modélisation UML (diagrammes de classes et de séquence).\n- Création et structuration de la base de données relationnelle sous PostgreSQL.\nRésultat : Le modèle de base de données est prêt et les tables principales (utilisateurs, véhicules, réservations) sont créées.")
    
    # Mardi
    p = doc.add_paragraph()
    r = p.add_run("Mardi 18/08/2026")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 73, 125)
    doc.add_paragraph("- Initialisation du projet Backend avec Node.js et Express.\n- Configuration de l'ORM Sequelize pour la connexion à la base de données PostgreSQL.\nRésultat : L'environnement serveur est opérationnel et connecté à la base de données.")
    
    # Mercredi
    p = doc.add_paragraph()
    r = p.add_run("Mercredi 19/08/2026")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 73, 125)
    doc.add_paragraph("- Développement du système d'authentification et d'inscription (API).\n- Mise en place de la sécurité avec le hachage des mots de passe (bcrypt) et les tokens JWT.\nRésultat : Les utilisateurs et administrateurs peuvent s'inscrire et se connecter de manière sécurisée.")
    
    # Jeudi
    p = doc.add_paragraph()
    r = p.add_run("Jeudi 20/08/2026")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 73, 125)
    doc.add_paragraph("- Développement des API REST pour la gestion du catalogue (CRUD services et produits).\n- Développement des API pour la gestion de la flotte de véhicules.\nRésultat : Les routes API permettent d'ajouter, modifier et consulter les produits et véhicules de l'entreprise.")
    
    # Vendredi
    p = doc.add_paragraph()
    r = p.add_run("Vendredi 21/08/2026")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 73, 125)
    doc.add_paragraph("- Implémentation de la logique de réservation de véhicules avec vérification des disponibilités (algorithme anti-double-réservation).\n- Mise en place du système de gestion du panier.\nRésultat : Le système est capable d'enregistrer des réservations en évitant les conflits de dates.")
    
    # Samedi
    p = doc.add_paragraph()
    r = p.add_run("Samedi 22/08/2026")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 73, 125)
    doc.add_paragraph("- Intégration du module de génération de fichiers PDF pour les devis.\n- Configuration des envois d'e-mails transactionnels (notifications).\nRésultat : Le serveur génère automatiquement des devis au format PDF et notifie les clients par courriel.")
    
    # Résultats obtenus
    p = doc.add_paragraph()
    r = p.add_run("Résultats obtenus")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(31, 73, 125)
    doc.add_paragraph("L'architecture Backend (API REST) est désormais en grande partie fonctionnelle. La base de données PostgreSQL est opérationnelle et les principales routes métier (Authentification, Catalogue, Véhicules, Réservations, Devis) sont implémentées et sécurisées. Les fondations logicielles de la plateforme sont solides.")
    
    # Perspectives
    p = doc.add_paragraph()
    r = p.add_run("Perspectives")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(31, 73, 125)
    doc.add_paragraph("- Finaliser les tests d'intégration des API (avec Postman).\n- Débuter le développement de l'interface Frontend Web (React.js) pour consommer ces API.\n- Corriger les éventuels bugs remontés lors des tests.")
    
    # Signature
    doc.add_paragraph()
    doc.add_paragraph()
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sig = p_sig.add_run("SORHO DAVID PAUL")
    r_sig.bold = True
    
    doc.save("c:\\Users\\HP\\Downloads\\PARCOURS TS\\STAGE\\STAGE TS STIC 2\\site-soutarah\\Rapport_Semaine_17_22_Aout.docx")
    print("Report generated successfully.")

if __name__ == "__main__":
    build_weekly_report()
