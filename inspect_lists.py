import docx

doc = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

with open('inspect_lists.txt', 'w', encoding='utf-8') as f:
    f.write("=== SOMMAIRE (p. 29 to 53) ===\n")
    for i in range(29, 54):
        f.write(f"{i}: {doc.paragraphs[i].text}\n")
        
    f.write("\n=== LISTE DES FIGURES (p. 54 to 84) ===\n")
    for i in range(54, 85):
        f.write(f"{i}: {doc.paragraphs[i].text}\n")
        
    f.write("\n=== LISTE DES TABLEAUX (p. 85 to 91) ===\n")
    for i in range(85, 92):
        f.write(f"{i}: {doc.paragraphs[i].text}\n")
        
    f.write("\n=== TABLE DES MATIERES (p. 710 to end) ===\n")
    for i in range(710, len(doc.paragraphs)):
        f.write(f"{i}: {doc.paragraphs[i].text}\n")

print("Lists written to inspect_lists.txt")
