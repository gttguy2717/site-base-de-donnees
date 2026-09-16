import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pres_common import (
    LOGO_DARK_GREEN, LOGO_BRAND_GREEN, LOGO_LEAF_LIME, LIME_LIGHT_BG, MINT_LIGHT_BG,
    DARK_CARD_BG, DARK_CARD_BORDER, LIGHT_BG, CARD_BG, CARD_BORDER,
    TEXT_DARK, TEXT_BODY, TEXT_MUTED, TEXT_WHITE,
    FONT_HEADING, FONT_BODY, set_slide_background, add_header, add_footer,
    add_card, add_badge, create_section_divider
)

def build_part5(prs):
    blank_layout = prs.slide_layouts[6]

    # ==========================================
    # SLIDE 49 : INTERCALAIRE CONCLUSION
    # ==========================================
    create_section_divider(prs, 49, 5,
                           "CONCLUSION GÉNÉRALE & PERSPECTIVES",
                           "Bilan du stage d'application, synthèse des résultats obtenus, compétences informatiques consolidées et perspectives de déploiement commercial.")

    # ==========================================
    # SLIDE 50 : CONCLUSION GÉNÉRALE & REMERCIEMENTS (SOUTARAH GREEN HERO)
    # ==========================================
    s50 = prs.slides.add_slide(blank_layout)
    set_slide_background(s50, prs, LOGO_DARK_GREEN)

    # Accent bar - Vert feuille signature
    top_bar = s50.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = LOGO_LEAF_LIME
    top_bar.line.fill.background()

    add_badge(s50, Inches(0.8), Inches(0.6), Inches(4.2), Inches(0.35),
              "BILAN DE STAGE • CYCLE TECHNICIEN SUPÉRIEUR",
              bg_color=DARK_CARD_BG, txt_color=LOGO_LEAF_LIME)

    tb_t = s50.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.73), Inches(0.6))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p1 = tf_t.paragraphs[0]
    p1.text = "Conclusion Générale & Remerciements au Jury"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

    concl_cards = [
        ("OBJECTIFS 100% ATTEINTS", "Une Solution Logicielle Prête pour la Production",
         "Conception et mise en œuvre réussie d'un écosystème complet : portail Web responsive, application mobile dédiée à la réservation de véhicules et back-office de gestion interconnectés en temps réel avec génération automatique de devis PDF pro-forma."),
        ("ACQUIS PROFESSIONNELS", "Compétences d'Ingénierie Consolidées",
         "Maîtrise approfondie de la méthode PU/UML pour la modélisation formelle, expertise sur la pile Full Stack moderne (React, React Native, Node.js, Express, MySQL Sequelize), rigueur de gestion de projet et développement de l'autonomie professionnelle."),
        ("REMERCIEMENTS RESPECTUEUX", "Gratitude aux Encadrants et au Jury",
         "Sincère gratitude aux éminents membres du jury de soutenance, à la direction générale et aux équipes de SOUTARAH GROUP pour leur accueil et leur confiance, ainsi qu'au corps professoral et à la direction de l'INP-HB / ESI / STIC.")
    ]
    w_c = Inches(3.75)
    h_c = Inches(4.3)
    for idx, (tag, title, desc) in enumerate(concl_cards):
        cx = Inches(0.8 + idx * 3.99)
        cy = Inches(1.95)
        top_c = LOGO_LEAF_LIME if idx==2 else LOGO_BRAND_GREEN
        add_card(s50, cx, cy, w_c, h_c, bg_color=DARK_CARD_BG, border_color=DARK_CARD_BORDER, top_accent_color=top_c)
        add_badge(s50, cx + Inches(0.3), cy + Inches(0.3), Inches(3.1), Inches(0.32), tag,
                  bg_color=LOGO_DARK_GREEN, txt_color=top_c)
        
        tb = s50.shapes.add_textbox(cx + Inches(0.3), cy + Inches(0.85), w_c - Inches(0.6), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_WHITE
        
        pd = tf.add_paragraph()
        pd.space_before = Pt(10)
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(11)
        pd.font.color.rgb = RGBColor(203, 213, 225)

    # Bottom Contact Bar
    add_badge(s50, Inches(2.5), Inches(6.5), Inches(8.33), Inches(0.4),
              "DAVID SORHO • ÉLÈVE TECHNICIEN SUPÉRIEUR • INP-HB YAMOUSSOUKRO / ESI / STIC",
              bg_color=DARK_CARD_BG, txt_color=LOGO_LEAF_LIME)

    add_footer(s50, 50, is_dark=True)
    print("Part 5 (Slides 49 and 50) updated with Soutarah Green theme.")
