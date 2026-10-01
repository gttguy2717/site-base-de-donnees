import docx
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc_fr = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')
with open('complete_translations.json', encoding='utf-8') as f:
    ct = json.load(f)

print(f"Total FR paragraphs: {len(doc_fr.paragraphs)}")

# Let's inspect each paragraph:
for i, p in enumerate(doc_fr.paragraphs):
    fr = p.text.strip()
    if not fr:
        continue
    en = ct.get(str(i), '').strip()
    
    # Simple semantic similarity / keyword check:
    # Numbers at start:
    fr_words = fr.split()
    en_words = en.split()
    fr_start = fr_words[0] if fr_words else ""
    en_start = en_words[0] if en_words else ""
    
    # Check if heading numbering matches:
    if fr_start in ['1.', '2.', '3.', '4.', '5.', '6.', 'I.', 'II.', 'III.', 'IV.', 'A.', 'B.', 'C.', 'D.']:
        if fr_start != en_start:
            print(f"SHIFT/MISMATCH at P[{i}]:")
            print(f"   FR: {repr(fr[:80])}")
            print(f"   EN: {repr(en[:80])}")
            # check neighbor EN
            for offset in [-2, -1, 1, 2]:
                if str(i+offset) in ct:
                    print(f"   EN[{i+offset}]: {repr(ct[str(i+offset)][:60])}")
            print("-" * 60)
