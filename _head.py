import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.oxml.ns import qn
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br"
p = r"C:\Users\HP\Downloads\PARCOURS TS\STAGE\STAGE TS STIC 2\rapport-stage\mon-rapport\Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx"
d = Document(p)
for i, par in enumerate(d.paragraphs[:60]):
    t = par.text.strip()
    img = len(par._p.findall(".//"+A))
    brk = any(b.get(W+"type")=="page" for b in par._p.findall(".//"+W))
    if t or img or brk:
        print(i, "|", par.style.name, "| img=%d"%img, "PAGEBREAK" if brk else "", "|", t[:110])
print()
print("--- sections ---")
for k,s in enumerate(d.sections):
    print("section",k,"start_type",s.start_type,"diff_first_page",s.different_first_page_header_footer)
    print("  header:", [x.text for x in s.header.paragraphs])
    print("  footer:", [x.text for x in s.footer.paragraphs])
