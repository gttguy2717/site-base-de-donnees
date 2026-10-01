# -*- coding: utf-8 -*-
"""
compare_fr_en.py
Compare French corrected doc vs English doc paragraph by paragraph.
Identifies mismatches in Genius Pay section and other key areas.
"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

FR = 'Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx'
EN = 'Rapport_de_Stage_ENGLISH.docx'

doc_fr = Document(FR)
doc_en = Document(EN)

# Print full French Genius Pay section text
print('=== GENIUS PAY - Texte FRANCAIS (para 555-563) ===')
for i in range(555, 564):
    p = doc_fr.paragraphs[i]
    print(f'[{i}][{p.style.name}] repr: {repr(p.text)}')

print()
print('=== GENIUS PAY - Texte ANGLAIS (para 541-552) ===')
for i in range(541, 553):
    if i >= len(doc_en.paragraphs): break
    p = doc_en.paragraphs[i]
    print(f'[{i}][{p.style.name}] repr: {repr(p.text)}')

# Also check the II section (Brevo + Genius Pay heading in French)
print()
print('=== FR: Section II heading (para 543) ===')
for i in range(543, 558):
    p = doc_fr.paragraphs[i]
    if p.text.strip():
        print(f'[{i}][{p.style.name}] {p.text[:120]}')

print()
print('=== EN: Section II (para 528-545) ===')
for i in range(528, 546):
    if i >= len(doc_en.paragraphs): break
    p = doc_en.paragraphs[i]
    if p.text.strip():
        print(f'[{i}][{p.style.name}] {p.text[:120]}')
