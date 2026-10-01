# -*- coding: utf-8 -*-
"""
dump_fr_text.py - Dump all French paragraphs for translation
"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

doc = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')
print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Total tables: {len(doc.tables)}')

with open('fr_paragraphs_dump.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        has_img = len(p._element.findall('.//' + qn('a:blip'))) > 0
        img_mark = '[IMG]' if has_img else ''
        if t:
            f.write(f'[{i}][{p.style.name}]{img_mark} {t}\n')

print('Dumped to fr_paragraphs_dump.txt')

# Also dump tables
with open('fr_tables_dump.txt', 'w', encoding='utf-8') as f:
    for t_idx, table in enumerate(doc.tables):
        f.write(f'\n=== TABLE {t_idx} ===\n')
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            f.write(' | '.join(cells) + '\n')

print('Tables dumped to fr_tables_dump.txt')
