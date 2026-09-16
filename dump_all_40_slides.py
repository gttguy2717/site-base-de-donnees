import pptx
import sys
sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('Soutenance stage plateforme numerique Soutarah Group.pptx')
print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    texts = []
    has_img = False
    for s in slide.shapes:
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            has_img = True
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                t = p.text.strip()
                if t and t not in ['David Sorho', '9/15/2026', f'{idx+1:02d}.', f'{idx+1}.']:
                    texts.append(t)
    
    title = texts[0] if texts else "SANS TITRE"
    body = " | ".join(texts[1:5]) if len(texts) > 1 else ""
    img_tag = "[IMAGE]" if has_img else "[NO-IMG]"
    print(f"Slide {idx+1:02d} {img_tag}: {title} ===> {body[:80]}")
