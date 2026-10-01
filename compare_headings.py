import sys
from docx import Document

doc_fr = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')
doc_en = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_ENGLISH.docx')

print("=== FR DOC SECTIONS & HEADINGS ===")
for i, p in enumerate(doc_fr.paragraphs):
    if p.style.name.startswith('Heading') and p.text.strip():
        print(f"FR [{i}] ({p.style.name}): {p.text.strip()[:60]}")

print("\n=== EN DOC SECTIONS & HEADINGS ===")
for i, p in enumerate(doc_en.paragraphs):
    if p.style.name.startswith('Heading') and p.text.strip():
        print(f"EN [{i}] ({p.style.name}): {p.text.strip()[:60]}")
