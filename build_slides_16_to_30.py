import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from build_common import (
    set_slide_background, add_header, add_footer, add_card, add_badge,
    create_section_divider, place_image_in_box, create_table,
    LOGO_DARK_GREEN, LOGO_BRAND_GREEN, LOGO_LEAF_LIME, LIME_LIGHT_BG, MINT_LIGHT_BG,
    DARK_CARD_BG, LIGHT_BG, CARD_BG, CARD_BORDER, TEXT_DARK, TEXT_BODY, TEXT_MUTED,
    TEXT_WHITE, FONT_HEADING, FONT_BODY
)

def build_slides_16_to_30(prs):
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 16 : SECTION 3 - PRÉSENTATION DE MERISE
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_slide_background(s16, prs, LIGHT_BG)
    add_header(s16, "3. MÉTHODE D’ANALYSE • APPROCHE TRADITIONNELLE",
               "Présentation de la Méthode MERISE",
               "Méthode cartésienne d'analyse et de conception des systèmes d'information")

    cards_merise = [
        ("PHILOSOPHIE & CYCLE", "Une Démarche Séquentielle",
         "• Méthode française de référence née dans les années 1970 pour les grands systèmes de gestion.\n• Approche en cycle de vie linéaire et séquentiel (cycle en V ou en cascade).\n• Séparation formelle et stricte entre l'univers des Données et l'univers des Traitements."),
        ("NIVEAUX D'ABSTRACTION", "Les Trois Niveaux d'Analyse",
         "• Niveau Conceptuel : MCD (Modèle Conceptuel de Données) et MCT (Modèle Conceptuel des Traitements).\n• Niveau Organisationnel : MLD (Modèle Logique de Données) et MOT (Modèle Organisationnel des Traitements).\n• Niveau Opérationnel : MPD (Modèle Physique de Données) et ULT (Unités Logiques de Traitement)."),
        ("FORCES & LIMITES", "Bilan Opérationnel",
         "• Grande robustesse pour la modélisation des bases de données relationnelles SQL classiques.\n• Forte rigidité : inadaptée aux changements fréquents d'exigences en cours de projet.\n• Rupture conceptuelle avec la programmation orientée objet moderne (React, Node.js).")
    ]

    for i, (b_txt, title, body) in enumerate(cards_merise):
        left = Inches(0.8) + i * Inches(3.95)
        top = Inches(1.85)
        width = Inches(3.8)
        height = Inches(4.9)
        add_card(s16, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
        add_badge(s16, left + Inches(0.25), top + Inches(0.25), Inches(3.0), Inches(0.32), b_txt)

        tb = s16.shapes.add_textbox(left + Inches(0.25), top + Inches(0.75), Inches(3.3), Inches(3.9))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13.5)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK

        pb = tf.add_paragraph()
        pb.text = body
        pb.font.name = FONT_BODY
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = TEXT_BODY
        pb.space_before = Pt(10)

    add_footer(s16, 16)

    # =========================================================================
    # SLIDE 17 : SECTION 3 - PRÉSENTATION DU PU ET D'UML
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    set_slide_background(s17, prs, LIGHT_BG)
    add_header(s17, "3. MÉTHODE D’ANALYSE • APPROCHE MODERNE",
               "Présentation du Processus Unifié (PU) et d'UML",
               "Cadre méthodologique itératif, incrémental et standardisé pour le génie logiciel")

    cards_pu = [
        ("PROCESSUS UNIFIÉ (PU)", "Les Piliers Méthodologiques",
         "• Piloté par les Cas d'Utilisation : l'utilisateur final et ses besoins réels guident chaque étape.\n• Centré sur l'Architecture : structuration robuste dès les premières itérations.\n• Itératif & Incrémental : décomposition du projet en sous-ensembles fonctionnels testés en continu.\n• Découpage en 4 phases : Cadrage, Élaboration, Construction, Transition."),
        ("LANGAGE UML 2.5", "Standard de Modélisation",
         "• Unified Modeling Language (UML) : langage graphique universel standardisé par l'OMG.\n• Vue Fonctionnelle : Diagramme de Cas d'Utilisation (Use Cases).\n• Vue Dynamique / Comportementale : Diagrammes d'Activité et de Séquence (flux opérationnels).\n• Vue Statique : Diagramme de Classes reflétant les entités de l'ORM Sequelize."),
        ("VALEUR AJOUTÉE", "Bénéfices pour le Projet",
         "• Continuité conceptuelle parfaite avec les architectures full-stack JavaScript / TypeScript.\n• Adaptabilité remarquable aux retours utilisateurs et aux ajustements commerciaux rapides.\n• Documentation vivante, lisible par l'équipe technique et la direction de SOUTARAH GROUP.")
    ]

    for i, (b_txt, title, body) in enumerate(cards_pu):
        left = Inches(0.8) + i * Inches(3.95)
        top = Inches(1.85)
        width = Inches(3.8)
        height = Inches(4.9)
        add_card(s17, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
        add_badge(s17, left + Inches(0.25), top + Inches(0.25), Inches(3.2), Inches(0.32), b_txt)

        tb = s17.shapes.add_textbox(left + Inches(0.25), top + Inches(0.75), Inches(3.3), Inches(3.9))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13.5)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK

        pb = tf.add_paragraph()
        pb.text = body
        pb.font.name = FONT_BODY
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = TEXT_BODY
        pb.space_before = Pt(10)

    add_footer(s17, 17)

    # =========================================================================
    # SLIDE 18 : SECTION 3 - ÉTUDE COMPARATIVE MERISE vs PU / UML
    # =========================================================================
    s18 = prs.slides.add_slide(blank_layout)
    set_slide_background(s18, prs, LIGHT_BG)
    add_header(s18, "3. MÉTHODE D’ANALYSE • COMPARAISON APPROFONDIE",
               "Étude Comparative : MERISE vs PU / UML (Tableau 2)",
               "Évaluation multicritère des deux approches méthodologiques d'ingénierie")

    headers_comp = ["Critères d'évaluation", "Méthode MERISE", "Processus Unifié & UML (PU/UML)"]
    rows_comp = [
        ["Approche conceptuelle", "Séquentielle (cycle en V), séparation stricte données et traitements", "Orientée objet, itérative, incrémentale et centrée sur l'utilisateur"],
        ["Modélisation des données", "MCD, MLD très adaptés aux bases de données SQL relationnelles", "Diagramme de classes reflétant directement les modèles d'objets (ORM)"],
        ["Modélisation des traitements", "MCT, MLT axés sur la circulation formelle des flux d'informations", "Diagrammes de cas d'utilisation, de séquence et d'activité interactifs"],
        ["Adaptation Web & Mobile", "Rigide pour les applications temps réel, asynchrones et événementielles", "Parfaite adéquation avec les frameworks modernes (React, Node.js, REST)"],
        ["Évolution & Agilité", "Reprise lourde et coûteuse des modèles conceptuels en cas de changement", "Intégration fluide de nouvelles fonctionnalités par cycles d'itérations courts"],
        ["Décision retenue", "Utilisée à titre informatif pour consolider la structure relationnelle", "Méthode principale retenue pour l'ingénierie globale de la solution"]
    ]
    col_w_comp = [Inches(2.73), Inches(4.5), Inches(4.5)]
    create_table(s18, Inches(0.8), Inches(1.85), Inches(11.73), Inches(4.8), headers_comp, rows_comp, col_w_comp)

    add_footer(s18, 18)

    # =========================================================================
    # SLIDE 19 : SECTION 3 - CHOIX DE LA MÉTHODE
    # =========================================================================
    s19 = prs.slides.add_slide(blank_layout)
    set_slide_background(s19, prs, LIGHT_BG)
    add_header(s19, "3. MÉTHODE D’ANALYSE • JUSTIFICATION DU CHOIX",
               "Choix de la Méthode : Pourquoi le PU/UML ?",
               "Une démarche en phase avec les impératifs d'agilité et de robustesse de SOUTARAH GROUP")

    raisons = [
        ("SYNERGIE ORIENTÉE OBJET", "Cohérence Full-Stack JS",
         "La modélisation UML s'accorde naturellement avec la programmation orientée composants de React, la logique modulaire d'Express et les modèles objet de l'ORM Sequelize."),
        ("CENTRÉE CAS D'UTILISATION", "Adéquation aux Besoins Réels",
         "Chaque cas d'utilisation (réserver un véhicule, demander un devis, payer en ligne) a directement guidé le développement des interfaces et des routes de l'API REST."),
        ("DÉMARCHE ITÉRATIVE RAPIDE", "Flexibilité face aux Évolutions",
         "Dans un stage court de deux mois, l'approche incrémentale a permis de livrer et tester des modules fonctionnels chaque quinzaine (Auth, Catalogue, Réservations, Devis)."),
        ("COMPRÉHENSION PARTAGÉE", "Communication Technique & Métier",
         "Les diagrammes UML constituent un langage visuel clair, facilitant la validation des flux opérationnels avec l'équipe dirigeante de SOUTARAH GROUP.")
    ]

    for idx, (badge_t, title, desc) in enumerate(raisons):
        col = 0 if idx < 2 else 1
        row = idx % 2
        left = Inches(0.8) if col == 0 else Inches(6.8)
        top = Inches(1.85) + row * Inches(2.45)
        width = Inches(5.7)
        height = Inches(2.25)

        add_card(s19, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
        add_badge(s19, left + Inches(0.25), top + Inches(0.25), Inches(3.2), Inches(0.32), badge_t)

        tb = s19.shapes.add_textbox(left + Inches(0.25), top + Inches(0.7), Inches(5.2), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_BODY
        pd.space_before = Pt(6)

    add_footer(s19, 19)

    # =========================================================================
    # SLIDE 20 : INTERCALAIRE SECTION 4 : ANALYSE ET SPÉCIFICATION
    # =========================================================================
    create_section_divider(prs, 20, 4, "ANALYSE ET SPÉCIFICATION DES BESOINS",
                           "Identification des acteurs du système, formalisation des besoins fonctionnels et non fonctionnels, et cartographie par les diagrammes de cas d'utilisation.")

    # =========================================================================
    # SLIDE 21 : SECTION 4 - IDENTIFICATION DES ACTEURS
    # =========================================================================
    s21 = prs.slides.add_slide(blank_layout)
    set_slide_background(s21, prs, LIGHT_BG)
    add_header(s21, "4. ANALYSE ET SPÉCIFICATION • ACTEURS",
               "Identification et Rôles des Acteurs du Système",
               "Caractérisation des entités humaines et logicielles interagissant avec la plateforme")

    acteurs = [
        ("ACTEUR PRIMAIRE", "Visiteur (Non Authentifié)",
         "• Rôle : Découverte de l'entreprise et de son offre.\n• Actions autorisées :\n  - Consulter le portail vitrine et les 6 pôles.\n  - Parcourir le catalogue et la flotte de véhicules.\n  - Dialoguer avec l'assistant virtuel IA.\n  - Créer un compte client (Particulier ou Entreprise)."),
        ("ACTEUR PRIMAIRE", "Client Authentifié (Web & Mobile)",
         "• Rôle : Commanditaire de services et réservations.\n• Actions autorisées :\n  - Sélectionner des dates de location de véhicules.\n  - Composer un panier multi-articles (véhicules + négoce).\n  - Valider des demandes de devis et télécharger le PDF.\n  - Effectuer le paiement en ligne (Genius Pay / Mobile Money).\n  - Suivre l'historique et les statuts de traitement."),
        ("ACTEUR SECONDAIRE", "Administrateur (Back-Office)",
         "• Rôle : Pilotage commercial et opérationnel.\n• Actions autorisées :\n  - Consulter le tableau de bord consolidé et KPIs.\n  - Gérer la flotte des 21 véhicules et disponibilités.\n  - Mettre à jour les articles de négoce et tarifs.\n  - Traiter les demandes de devis et marquer comme lu.\n  - Superviser les clients et les notifications."),
        ("ACTEUR LOGICIEL", "Système / API REST & Tiers",
         "• Rôle : Orchestration des règles métier automatisées.\n• Actions autorisées :\n  - Contrôle algorithmique anti-collision des dates.\n  - Calcul automatique des taxes (TVA 18%, TDT 2.5%).\n  - Génération dynamique des devis PDF officiels 2 pages.\n  - Émission des e-mails transactionnels via Brevo.\n  - Sécurisation des routes par jetons JWT.")
    ]

    for idx, (b_t, title, desc) in enumerate(acteurs):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8) if col == 0 else Inches(6.8)
        top = Inches(1.85) + row * Inches(2.45)
        width = Inches(5.7)
        height = Inches(2.25)

        add_card(s21, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
        add_badge(s21, left + Inches(0.25), top + Inches(0.25), Inches(2.8), Inches(0.32), b_t)

        tb = s21.shapes.add_textbox(left + Inches(0.25), top + Inches(0.65), Inches(5.2), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = TEXT_BODY
        pd.space_before = Pt(4)

    add_footer(s21, 21)

    # =========================================================================
    # SLIDE 22 : SECTION 4 - BESOINS FONCTIONNELS
    # =========================================================================
    s22 = prs.slides.add_slide(blank_layout)
    set_slide_background(s22, prs, LIGHT_BG)
    add_header(s22, "4. ANALYSE ET SPÉCIFICATION • BESOINS FONCTIONNELS",
               "Spécification Détaillée des Besoins Fonctionnels",
               "Les exigences métiers indispensables déployées sur le Web et le Mobile")

    # 2 colonnes : Client vs Administrateur
    add_card(s22, Inches(0.8), Inches(1.85), Inches(5.7), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s22, Inches(1.05), Inches(2.1), Inches(3.2), Inches(0.32), "EXIGENCES CÔTÉ CLIENT (WEB & MOBILE)")

    tb_bf_c = s22.shapes.add_textbox(Inches(1.05), Inches(2.55), Inches(5.2), Inches(4.0))
    tf_bfc = tb_bf_c.text_frame
    tf_bfc.word_wrap = True
    p1 = tf_bfc.paragraphs[0]
    p1.text = "Fonctionnalités Disponibles pour l'Usager"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_DARK

    p1b = tf_bfc.add_paragraph()
    p1b.text = "• Consultation interactive du catalogue multi-services des 6 pôles.\n\n" \
               "• Réservation de véhicules : filtrage multicritère (SUV, Berline, boîte auto/manuelle), sélection calendrier, option chauffeur et destination.\n\n" \
               "• Panier multi-articles hybride : combinaison de locations de véhicules et de produits matériels dans un panier unique.\n\n" \
               "• Émission de devis : génération et téléchargement direct du devis pro-forma officiel PDF conforme à la charte de l'entreprise.\n\n" \
               "• Espace personnel : suivi en direct des statuts de devis (En attente, Lu, Confirmé) et notifications reçues."
    p1b.font.name = FONT_BODY
    p1b.font.size = Pt(10.5)
    p1b.font.color.rgb = TEXT_BODY
    p1b.space_before = Pt(8)

    add_card(s22, Inches(6.8), Inches(1.85), Inches(5.7), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s22, Inches(7.05), Inches(2.1), Inches(3.2), Inches(0.32), "EXIGENCES CÔTÉ ADMIN (BACK-OFFICE)")

    tb_bf_a = s22.shapes.add_textbox(Inches(7.05), Inches(2.55), Inches(5.2), Inches(4.0))
    tf_bfa = tb_bf_a.text_frame
    tf_bfa.word_wrap = True
    p2 = tf_bfa.paragraphs[0]
    p2.text = "Fonctionnalités de Pilotage Administrateur"
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_DARK

    p2b = tf_bfa.add_paragraph()
    p2b.text = "• Tableau de bord dynamique : métriques en temps réel (chiffre d'affaires estimé, volume de devis, réservations actives).\n\n" \
               "• Gestion de la flotte automobile : ajout, modification technique, tarif journalier, statut de disponibilité des 21 véhicules.\n\n" \
               "• Gestion des produits de négoce : actualisation du catalogue, photos, prix unitaires et fiches techniques.\n\n" \
               "• Traitement des devis : consultation rapide, action « Marquer comme lu » pour enclencher la préparation sans blocage.\n\n" \
               "• Gestion de la relation client : suivi des comptes particuliers/entreprises et historique complet des commandes."
    p2b.font.name = FONT_BODY
    p2b.font.size = Pt(10.5)
    p2b.font.color.rgb = TEXT_BODY
    p2b.space_before = Pt(8)

    add_footer(s22, 22)

    # =========================================================================
    # SLIDE 23 : SECTION 4 - BESOINS NON FONCTIONNELS
    # =========================================================================
    s23 = prs.slides.add_slide(blank_layout)
    set_slide_background(s23, prs, LIGHT_BG)
    add_header(s23, "4. ANALYSE ET SPÉCIFICATION • BESOINS NON FONCTIONNELS",
               "Besoins Non Fonctionnels & Contraintes Techniques",
               "Les exigences de qualité, de sécurité et de robustesse garantissant la pérennité")

    bnf = [
        ("SÉCURITÉ APPLICATIVE", "Authentification & Chiffrement",
         "• Authentification sans état via JSON Web Token (JWT) avec expiration sécurisée.\n• Hachage robuste des mots de passe utilisateurs via Bcrypt (salt à 10 tours).\n• Protection des routes API sensibles par middleware de contrôle des rôles (CLIENT/ADMIN)."),
        ("FIABILITÉ & CONFLITS", "Intégrité & Contrôle d'Agenda",
         "• Vérification algorithmique systématique des disponibilités anti-double réservation.\n• Transactions SQL sécurisées garantissant la cohérence atomique des commandes (ACID).\n• Verrouillage des plages horaires pour empêcher les collisions de réservations concurrentes."),
        ("PERFORMANCE & SCALABILITÉ", "Temps de Réponse & Charge",
         "• Architecture REST optimisée échangeant des charges utiles JSON ultra-légères.\n• Temps de réponse des requêtes API standard inférieur à 200 ms.\n• Prise en charge fluide de requêtes simultanées sur l'application mobile et le portail web."),
        ("ERGONOMIE & NORMES", "Expérience Utilisateur & Formats",
         "• Interface fully responsive s'adaptant parfaitement aux écrans desktop, tablettes et smartphones.\n• Respect des normes locales : dates au format français (jj/mm/aaaa) et devises en FCFA (XOF).\n• Respect scrupuleux de la charte juridique et fiscale ivoirienne sur les devis (TVA 18%, TDT 2.5%).")
    ]

    for idx, (badge_t, title, desc) in enumerate(bnf):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8) if col == 0 else Inches(6.8)
        top = Inches(1.85) + row * Inches(2.45)
        width = Inches(5.7)
        height = Inches(2.25)

        add_card(s23, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
        add_badge(s23, left + Inches(0.25), top + Inches(0.25), Inches(3.2), Inches(0.32), badge_t)

        tb = s23.shapes.add_textbox(left + Inches(0.25), top + Inches(0.65), Inches(5.2), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_BODY
        pd.space_before = Pt(4)

    add_footer(s23, 23)

    # =========================================================================
    # SLIDE 24 : SECTION 4 - DIAGRAMME GLOBAL DES CAS D'UTILISATION (FIGURE 2)
    # =========================================================================
    s24 = prs.slides.add_slide(blank_layout)
    set_slide_background(s24, prs, LIGHT_BG)
    add_header(s24, "4. ANALYSE ET SPÉCIFICATION • MODÉLISATION FONCTIONNELLE",
               "Diagramme Global des Cas d'Utilisation (Figure 2)",
               "Cartographie des fonctionnalités de la plateforme SOUTARAH selon les profils d'acteurs")

    img_uc = "extracted_docx_images/image2.png"
    add_card(s24, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s24, img_uc, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    # Carte analyse droite
    add_card(s24, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s24, Inches(9.25), Inches(2.1), Inches(2.5), Inches(0.32), "ANALYSE DU DIAGRAMME")

    tb_uc = s24.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_uc = tb_uc.text_frame
    tf_uc.word_wrap = True
    p_uc = tf_uc.paragraphs[0]
    p_uc.text = "Synthèse des Interactions"
    p_uc.font.name = FONT_HEADING
    p_uc.font.size = Pt(13)
    p_uc.font.bold = True
    p_uc.font.color.rgb = TEXT_DARK

    p_ucb = tf_uc.add_paragraph()
    p_ucb.text = "• Visiteur : consultation libre des services, du catalogue, échange avec l'IA et inscription.\n\n" \
                 "• Client : réservation de véhicules, panier multi-articles, validation de devis, paiement en ligne et suivi.\n\n" \
                 "• Administrateur : gestion du parc automobile, validation/lecture des devis, actualisation catalogue et tableau de bord.\n\n" \
                 "• Relations dynamiques : 'include' pour la vérification de disponibilité et 'extend' pour les options avec chauffeur."
    p_ucb.font.name = FONT_BODY
    p_ucb.font.size = Pt(10)
    p_ucb.font.color.rgb = TEXT_BODY
    p_ucb.space_before = Pt(8)

    add_footer(s24, 24)

    # =========================================================================
    # SLIDE 25 : SECTION 4 - SCÉNARIOS DES CAS D'UTILISATION CLÉS
    # =========================================================================
    s25 = prs.slides.add_slide(blank_layout)
    set_slide_background(s25, prs, LIGHT_BG)
    add_header(s25, "4. ANALYSE ET SPÉCIFICATION • CAS D'UTILISATION CLÉS",
               "Description Textuelle des Cas d'Utilisation Majeurs",
               "Scénarios nominaux d'exécution pour les fonctionnalités critiques du système")

    uc_cases = [
        ("UC1 : RÉSERVER UN VÉHICULE", "Client Web / Mobile",
         "• Précondition : Client authentifié.\n• Scénario nominal :\n  1. Le client sélectionne un véhicule de la flotte.\n  2. Il choisit les dates de début et fin de location.\n  3. Le système vérifie la disponibilité sans conflit.\n  4. Calcul immédiat du tarif journalier et options.\n  5. Validation et ajout de la réservation au panier."),
        ("UC2 : DEMANDER UN DEVIS", "Client Authentifié",
         "• Précondition : Panier constitué (véhicules/produits).\n• Scénario nominal :\n  1. Le client accède au récapitulatif du panier.\n  2. Le système calcule les montants HT, TVA et TDT.\n  3. Le client confirme sa demande chiffrée.\n  4. Génération automatique du devis PDF officiel.\n  5. Téléchargement direct et transmission à l'admin."),
        ("UC3 : SUPERVISER L'ACTIVITÉ", "Administrateur Back-Office",
         "• Précondition : Compte administrateur connecté.\n• Scénario nominal :\n  1. L'administrateur accède au tableau de bord.\n  2. Consultation des indicateurs clés (CA, réservations).\n  3. Consultation de la liste des devis reçus.\n  4. Clic sur « Marquer comme lu » pour prise en compte.\n  5. Préparation logistique du véhicule ou colis.")
    ]

    for i, (b_t, title, body) in enumerate(uc_cases):
        left = Inches(0.8) + i * Inches(3.95)
        top = Inches(1.85)
        width = Inches(3.8)
        height = Inches(4.9)
        add_card(s25, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
        add_badge(s25, left + Inches(0.25), top + Inches(0.25), Inches(3.2), Inches(0.32), b_t)

        tb = s25.shapes.add_textbox(left + Inches(0.25), top + Inches(0.75), Inches(3.3), Inches(3.9))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13.5)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK

        pb = tf.add_paragraph()
        pb.text = body
        pb.font.name = FONT_BODY
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = TEXT_BODY
        pb.space_before = Pt(10)

    add_footer(s25, 25)

    # =========================================================================
    # SLIDE 26 : INTERCALAIRE SECTION 5 : CONCEPTION DU SYSTÈME
    # =========================================================================
    create_section_divider(prs, 26, 5, "CONCEPTION DU SYSTÈME",
                           "Modélisation dynamique par diagrammes d'activités et de séquences, structure statique des classes, architecture 3-tiers et choix technologiques.")

    # =========================================================================
    # SLIDE 27 : SECTION 5 - DIAGRAMME D'ACTIVITÉS 1 : RÉSERVATION VÉHICULE (FIGURE 4)
    # =========================================================================
    s27 = prs.slides.add_slide(blank_layout)
    set_slide_background(s27, prs, LIGHT_BG)
    add_header(s27, "5. CONCEPTION DU SYSTÈME • ACTIVITÉ VÉHICULE",
               "Processus de Réservation de Véhicule (Figure 4)",
               "Contrôle algorithmique anti-collision des dates et traitement simplifié côté administration")

    img_act_v = "extracted_docx_images/image4.png"
    add_card(s27, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s27, img_act_v, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    add_card(s27, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s27, Inches(9.25), Inches(2.1), Inches(2.8), Inches(0.32), "ANALYSE DU FLUX MÉTIER")

    tb_av = s27.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_av = tb_av.text_frame
    tf_av.word_wrap = True
    pav = tf_av.paragraphs[0]
    pav.text = "Étapes du Processus"
    pav.font.name = FONT_HEADING
    pav.font.size = Pt(13)
    pav.font.bold = True
    pav.font.color.rgb = TEXT_DARK

    pav_b = tf_av.add_paragraph()
    pav_b.text = "• Sélection du véhicule et saisie des dates de début et de fin de location.\n\n" \
                 "• Test de disponibilité en base de données :\n" \
                 "  - Si conflit : alerte immédiate au client pour réajuster ses dates.\n" \
                 "  - Si libre : ajout fluide au panier.\n\n" \
                 "• Validation du panier ➜ génération automatique du devis pro-forma.\n\n" \
                 "• Notification instantanée de l'administrateur qui marque simplement comme lu pour préparer le véhicule."
    pav_b.font.name = FONT_BODY
    pav_b.font.size = Pt(10)
    pav_b.font.color.rgb = TEXT_BODY
    pav_b.space_before = Pt(8)

    add_footer(s27, 27)

    # =========================================================================
    # SLIDE 28 : SECTION 5 - DIAGRAMME D'ACTIVITÉS 2 : ACHAT DE PRODUIT (FIGURE 5)
    # =========================================================================
    s28 = prs.slides.add_slide(blank_layout)
    set_slide_background(s28, prs, LIGHT_BG)
    add_header(s28, "5. CONCEPTION DU SYSTÈME • ACTIVITÉ ACHAT",
               "Processus d'Achat de Produits & Panier (Figure 5)",
               "Cinématique de commande directe de matériel de négoce et transparence des calculs")

    img_act_p = "extracted_docx_images/image5.png"
    add_card(s28, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s28, img_act_p, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    add_card(s28, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s28, Inches(9.25), Inches(2.1), Inches(2.8), Inches(0.32), "ANALYSE DU FLUX D'ACHAT")

    tb_ap = s28.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_ap = tb_ap.text_frame
    tf_ap.word_wrap = True
    pap = tf_ap.paragraphs[0]
    pap.text = "Cinématique Transactionnelle"
    pap.font.name = FONT_HEADING
    pap.font.size = Pt(13)
    pap.font.bold = True
    pap.font.color.rgb = TEXT_DARK

    pap_b = tf_ap.add_paragraph()
    pap_b.text = "• Parcours du catalogue négoce et ajout d'articles au panier.\n\n" \
                 "• Moteur de tarification : affichage transparent des montants Hors Taxes, TVA (18%) et Total TTC.\n\n" \
                 "• Choix du mode de règlement : en ligne (Genius Pay / Mobile Money) ou en espèces à la livraison.\n\n" \
                 "• Enregistrement de la commande et validation automatique du devis.\n\n" \
                 "• Notification admin pour ordonnancement du colisage et expédition."
    pap_b.font.name = FONT_BODY
    pap_b.font.size = Pt(10)
    pap_b.font.color.rgb = TEXT_BODY
    pap_b.space_before = Pt(8)

    add_footer(s28, 28)

    # =========================================================================
    # SLIDE 29 : SECTION 5 - DIAGRAMME D'ACTIVITÉS 3 : COMPTE CLIENT (FIGURE 3)
    # =========================================================================
    s29 = prs.slides.add_slide(blank_layout)
    set_slide_background(s29, prs, LIGHT_BG)
    add_header(s29, "5. CONCEPTION DU SYSTÈME • ACTIVITÉ UTILISATEUR",
               "Création et Gestion du Compte Client (Figure 3)",
               "Processus d'enrôlement sécurisé, vérification d'unicité et confirmation par email")

    img_act_c = "extracted_docx_images/image3.png"
    add_card(s29, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s29, img_act_c, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    add_card(s29, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s29, Inches(9.25), Inches(2.1), Inches(2.8), Inches(0.32), "ANALYSE DE L'ENRÔLEMENT")

    tb_ac = s29.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_ac = tb_ac.text_frame
    tf_ac.word_wrap = True
    pac = tf_ac.paragraphs[0]
    pac.text = "Parcours d'Inscription"
    pac.font.name = FONT_HEADING
    pac.font.size = Pt(13)
    pac.font.bold = True
    pac.font.color.rgb = TEXT_DARK

    pac_b = tf_ac.add_paragraph()
    pac_b.text = "• Saisie des informations client : sélection profil Particulier ou Entreprise (RCCM, NIF).\n\n" \
                 "• Contrôle d'intégrité API : vérification de l'unicité de l'e-mail et du numéro de téléphone.\n\n" \
                 "• Sécurité des données : hachage Bcrypt du mot de passe avant persistance en base MySQL.\n\n" \
                 "• Déclenchement automatique d'un courriel de bienvenue via l'API Brevo SMTP.\n\n" \
                 "• Redirection immédiate vers la page de connexion pour ouverture de session JWT."
    pac_b.font.name = FONT_BODY
    pac_b.font.size = Pt(10)
    pac_b.font.color.rgb = TEXT_BODY
    pac_b.space_before = Pt(8)

    add_footer(s29, 29)

    # =========================================================================
    # SLIDE 30 : SECTION 5 - DIAGRAMME DE SÉQUENCE 1 : AUTHENTIFICATION (FIGURE 9)
    # =========================================================================
    s30 = prs.slides.add_slide(blank_layout)
    set_slide_background(s30, prs, LIGHT_BG)
    add_header(s30, "5. CONCEPTION DU SYSTÈME • SÉQUENCE AUTH",
               "Authentification Sécurisée & Émission de Token JWT (Figure 9)",
               "Cinématique d'échange asynchrone entre Frontend, API REST, Base MySQL et Brevo")

    img_seq_auth = "extracted_docx_images/image9.png"
    add_card(s30, Inches(0.8), Inches(1.85), Inches(8.0), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER)
    place_image_in_box(s30, img_seq_auth, Inches(1.0), Inches(2.05), Inches(7.6), Inches(4.5))

    add_card(s30, Inches(9.0), Inches(1.85), Inches(3.53), Inches(4.9), bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s30, Inches(9.25), Inches(2.1), Inches(2.8), Inches(0.32), "ANALYSE DE SÉQUENCE")

    tb_sa = s30.shapes.add_textbox(Inches(9.25), Inches(2.6), Inches(3.03), Inches(4.0))
    tf_sa = tb_sa.text_frame
    tf_sa.word_wrap = True
    psa = tf_sa.paragraphs[0]
    psa.text = "Protocole d'Échange"
    psa.font.name = FONT_HEADING
    psa.font.size = Pt(13)
    psa.font.bold = True
    psa.font.color.rgb = TEXT_DARK

    psa_b = tf_sa.add_paragraph()
    psa_b.text = "1. Le formulaire Frontend transmet les credentials chiffrés à la route POST /api/auth/login.\n\n" \
                 "2. Le contrôleur API interroge la table Users et compare l'empreinte avec bcrypt.compare().\n\n" \
                 "3. Si valide, l'API génère un JWT signé contenant l'ID utilisateur et son rôle (CLIENT / ADMIN).\n\n" \
                 "4. Le Frontend stocke le token de session et déverrouille l'accès aux fonctionnalités protégées."
    psa_b.font.name = FONT_BODY
    psa_b.font.size = Pt(10)
    psa_b.font.color.rgb = TEXT_BODY
    psa_b.space_before = Pt(8)

    add_footer(s30, 30)

    print("Slides 16 to 30 successfully built.")
