import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

for i in [0, 1, 5, 6, 31, 54, 95, 105, 130, 210, 244, 250, 421, 510]:
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        print(f"=== Paragraph {i} (style={p.style.name}, runs={len(p.runs)}) ===")
        print("Text:", p.text[:80])
        for j, r in enumerate(p.runs):
            d = 'w:drawing' in r._r.xml
            print(f"  Run {j}: bold={r.bold}, italic={r.italic}, color={r.font.color.rgb if r.font.color else None}, drawing={d}, text={repr(r.text[:40])}")
