import os
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

# --- CHARTE OFFICIELLE SOUTARAH GROUP (100% VERT DU LOGO OFFICIEL) ---
LOGO_DARK_GREEN   = RGBColor(8, 44, 27)     # #082C1B - Vert forêt impérial profond (fonds sombres)
LOGO_BRAND_GREEN  = RGBColor(40, 129, 52)   # #288134 - Vert signature Soutarah (éléments principaux)
LOGO_LEAF_LIME    = RGBColor(120, 183, 43)  # #78B72B - Vert feuille / lime vif du logo (accents, kickers)
LIME_LIGHT_BG     = RGBColor(236, 248, 226) # #ECF8E2 - Vert très clair pour badges
MINT_LIGHT_BG     = RGBColor(220, 252, 231) # #DCFCE7 - Vert menthe doux
DARK_CARD_BG      = RGBColor(14, 60, 36)    # #0E3C24 - Carte sur fond sombre
DARK_CARD_BORDER  = RGBColor(40, 129, 52)   # #288134 - Bordure carte sombre

# Arrière-plans et cartes pour diapositives de contenu
LIGHT_BG          = RGBColor(248, 251, 248) # #F8FBF8 - Fond blanc végétal doux
CARD_BG           = RGBColor(255, 255, 255) # #FFFFFF - Blanc pur
CARD_BORDER       = RGBColor(226, 235, 228) # #E2EBE4 - Bordure subtile verdoyante

# Typographie moderne et élégante
TEXT_DARK         = RGBColor(15, 23, 42)    # #0F172A - Titres et texte noir ardoise
TEXT_BODY         = RGBColor(51, 65, 85)    # #334155 - Texte courant
TEXT_MUTED        = RGBColor(100, 116, 139) # #64748B - Métadonnées et sous-titres
TEXT_WHITE        = RGBColor(255, 255, 255) # Blanc
TABLE_HEADER_BG   = RGBColor(14, 60, 36)    # #0E3C24 - Vert forêt pour en-tête tableau
ROW_ALT_BG        = RGBColor(244, 249, 244) # #F4F9F4 - Alternance ligne tableau très claire

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"
TOTAL_SLIDES = 60

LOGO_PATH = "public/logo-soutarah.png"

def init_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def set_slide_background(slide, prs, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, kicker, title, subtitle=None, is_dark=False):
    # Kicker / Sur-titre
    tb_k = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.73), Inches(0.26))
    tf_k = tb_k.text_frame
    tf_k.word_wrap = True
    tf_k.margin_left = tf_k.margin_top = tf_k.margin_right = tf_k.margin_bottom = 0
    p_k = tf_k.paragraphs[0]
    p_k.text = kicker.upper()
    p_k.font.name = FONT_HEADING
    p_k.font.size = Pt(10)
    p_k.font.bold = True
    p_k.font.color.rgb = LOGO_LEAF_LIME if is_dark else LOGO_BRAND_GREEN

    # Grand Titre de Slide
    tb_t = slide.shapes.add_textbox(Inches(0.8), Inches(0.66), Inches(11.73), Inches(0.55))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(21)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE if is_dark else TEXT_DARK

    # Sous-titre si présent
    if subtitle:
        tb_s = slide.shapes.add_textbox(Inches(0.8), Inches(1.25), Inches(11.73), Inches(0.32))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = RGBColor(203, 213, 225) if is_dark else TEXT_MUTED

def add_footer(slide, slide_num, total_slides=TOTAL_SLIDES, is_dark=False):
    # Ligne de séparation discrète
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.73), Inches(0.012))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(25, 85, 50) if is_dark else RGBColor(220, 235, 225)
    line.line.fill.background()

    # Mention gauche
    tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(8.5), Inches(0.25))
    tf_l = tb_l.text_frame
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "SOUTARAH GROUP • Soutenance de Stage • David SORHO • INP-HB / ESI / STIC"
    p_l.font.name = FONT_BODY
    p_l.font.size = Pt(9)
    p_l.font.color.rgb = RGBColor(148, 163, 184) if is_dark else TEXT_MUTED

    # Numérotation droite
    tb_r = slide.shapes.add_textbox(Inches(10.53), Inches(7.12), Inches(2.0), Inches(0.25))
    tf_r = tb_r.text_frame
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.text = f"{slide_num:02d} / {total_slides:02d}"
    p_r.font.name = FONT_HEADING
    p_r.font.size = Pt(9.5)
    p_r.font.bold = True
    p_r.font.color.rgb = LOGO_LEAF_LIME if is_dark else LOGO_BRAND_GREEN

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER, top_accent_color=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    
    if top_accent_color:
        accent = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = top_accent_color
        accent.line.fill.background()
    return card

def add_badge(slide, left, top, width, height, text, bg_color=LIME_LIGHT_BG, txt_color=LOGO_BRAND_GREEN):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    pill.fill.solid()
    pill.fill.fore_color.rgb = bg_color
    pill.line.fill.background()
    tf = pill.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text.upper()
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_HEADING
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = txt_color
    return pill

def create_section_divider(prs, slide_num, seq_num, part_title, part_desc):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    set_slide_background(slide, prs, LOGO_DARK_GREEN)

    # Accent vert feuille logo
    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(0.12), Inches(3.2))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = LOGO_LEAF_LIME
    accent_bar.line.fill.background()

    add_badge(slide, Inches(1.4), Inches(2.2), Inches(2.2), Inches(0.36),
              f"SECTION {seq_num:02d}", bg_color=DARK_CARD_BG, txt_color=LOGO_LEAF_LIME)

    tb_t = slide.shapes.add_textbox(Inches(1.4), Inches(2.75), Inches(10.5), Inches(1.3))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = part_title
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE

    tb_d = slide.shapes.add_textbox(Inches(1.4), Inches(4.2), Inches(10.5), Inches(1.2))
    tf_d = tb_d.text_frame
    tf_d.word_wrap = True
    tf_d.margin_left = tf_d.margin_top = tf_d.margin_right = tf_d.margin_bottom = 0
    p_d = tf_d.paragraphs[0]
    p_d.text = part_desc
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(13.5)
    p_d.font.color.rgb = RGBColor(203, 213, 225)

    add_footer(slide, slide_num, is_dark=True)
    return slide

def place_image_in_box(slide, img_path, box_left, box_top, box_width, box_height):
    """Calcule le ratio d'aspect exact pour centrer l'image dans la boîte sans distorsion."""
    if not os.path.exists(img_path):
        print(f"Warning: image path does not exist: {img_path}")
        return None
    with Image.open(img_path) as im:
        img_w, img_h = im.size
    
    img_aspect = img_w / img_h
    box_aspect = box_width / box_height

    if img_aspect > box_aspect:
        final_w = box_width
        final_h = box_width / img_aspect
        final_left = box_left
        final_top = box_top + (box_height - final_h) / 2
    else:
        final_h = box_height
        final_w = box_height * img_aspect
        final_top = box_top
        final_left = box_left + (box_width - final_w) / 2

    return slide.shapes.add_picture(img_path, final_left, final_top, width=final_w, height=final_h)

def create_table(slide, left, top, width, height, headers, rows, col_widths=None):
    num_rows = len(rows) + 1
    num_cols = len(headers)
    tbl_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    tbl = tbl_shape.table

    if col_widths and len(col_widths) == num_cols:
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = w

    # Header row
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_HEADER_BG
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = str(h)
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.alignment = PP_ALIGN.LEFT

    # Data rows
    for i, row in enumerate(rows):
        bg = ROW_ALT_BG if i % 2 == 1 else CARD_BG
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = str(val)
            p.font.name = FONT_BODY
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_DARK
            p.alignment = PP_ALIGN.LEFT
    
    return tbl_shape
