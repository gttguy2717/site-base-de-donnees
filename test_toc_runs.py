from docx import Document

doc = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

def translate_toc_paragraph(para, english_title):
    """
    Translates a TOC line while strictly preserving its tab stop and page number.
    TOC line typically ends with '\t' and ' - X -' or similar.
    """
    runs = para.runs
    if not runs:
        return
    
    # Check if there is a tab in the runs
    tab_idx = None
    for j, r in enumerate(runs):
        if '\t' in r.text:
            tab_idx = j
            break
            
    if tab_idx is not None:
        # Title is in runs before tab_idx
        # Put english_title into the first text run before tab_idx
        first_title_run = False
        for j in range(tab_idx):
            if not first_title_run and runs[j].text.strip():
                runs[j].text = english_title
                first_title_run = True
            else:
                runs[j].text = ''
        # runs from tab_idx onwards (the tab and page number) are UNTOUCHED!
    else:
        # No tab found, simple replacement
        runs[0].text = english_title
        for r in runs[1:]:
            r.text = ''

# Test on P[30] (Sommaire DEDICACE) and P[55] (Figure 1)
translate_toc_paragraph(doc.paragraphs[30], "DEDICATION")
print("P[30] full text:", repr(doc.paragraphs[30].text))
print("P[30] runs:", [r.text for r in doc.paragraphs[30].runs])

translate_toc_paragraph(doc.paragraphs[55], "Figure 1 : GANTT Schedule Diagram")
print("\nP[55] full text:", repr(doc.paragraphs[55].text))
print("P[55] runs:", [r.text for r in doc.paragraphs[55].runs])
