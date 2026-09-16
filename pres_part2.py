import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pres_common import (
    LOGO_DARK_GREEN, LOGO_BRAND_GREEN, LOGO_LEAF_LIME, LIME_LIGHT_BG, MINT_LIGHT_BG,
    DARK_CARD_BG, DARK_CARD_BORDER, LIGHT_BG, CARD_BG, CARD_BORDER,
    TEXT_DARK, TEXT_BODY, TEXT_MUTED, TEXT_WHITE, TABLE_HEADER_BG, ROW_ALT_BG,
    FONT_HEADING, FONT_BODY, set_slide_background, add_header, add_footer,
    add_card, add_badge, create_section_divider, place_image_in_box
)

def build_part2(prs):
    blank_layout = prs.slide_layouts[6]

    # ==========================================
    # SLIDE 16 : INTERCALAIRE PARTIE II
    # ==========================================
    create_section_divider(prs, 16, 2,
                           "PARTIE II : ÉTUDE CONCEPTUELLE",
                           "Choix de la méthode d'analyse, justification de PU/UML, modélisation dynamique des cas d'utilisation, des activités et des séquences, et modélisation statique.")

    # ==========================================
    # SLIDE 17 : MERISE VS PU/UML (TABLEAU 3 DU RAPPORT)
    # ==========================================
    s17 = prs.slides.add_slide(blank_layout)
    set_slide_background(s17, prs, LIGHT_BG)
    add_header(s17, "Choix Méthodologique", "Analyse Comparative des Méthodes : MERISE vs PU/UML",
               "Évaluation rigoureuse selon 6 critères déterminants issue du rapport officiel (Tableau 3)")

    merise_data = [
        ("Critères d'évaluation", "Méthode MERISE", "Processus Unifié & UML (PU/UML)"),
        ("Approche", "Séquentielle, séparation stricte données / traitements", "Orientée objet, itérative, incrémentale et centrée usager"),
        ("Modélisation données", "MCD, MLD très adaptés aux bases SQL classiques", "Diagramme de classes reflétant directement les modèles ORM"),
        ("Modélisation traitements", "MCT, MLT axés sur les flux d'informations", "Diagrammes de cas d'utilisation, séquences et activités"),
        ("Architectures Web / Mobile", "Rigide pour les applications événementielles", "Parfaite adéquation avec les frameworks modernes (React, Node)"),
        ("Évolution du système", "Reprise lourde des modèles en cas de changement", "Intégration fluide de nouvelles fonctionnalités par itération"),
        ("Décision retenue", "Complément conceptuel pour la base de données", "Méthode principale retenue pour l'ingénierie globale")
    ]

    t_shape = s17.shapes.add_table(len(merise_data), 3, Inches(0.8), Inches(1.85), Inches(11.73), Inches(4.7))
    tbl = t_shape.table
    tbl.columns[0].width = Inches(2.2)
    tbl.columns[1].width = Inches(4.76)
    tbl.columns[2].width = Inches(4.76)

    for r_idx, row in enumerate(merise_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r_idx == 0:
                cell.fill.fore_color.rgb = TABLE_HEADER_BG
            elif r_idx == len(merise_data) - 1 and c_idx == 2:
                cell.fill.fore_color.rgb = LIME_LIGHT_BG
            elif r_idx % 2 == 1:
                cell.fill.fore_color.rgb = CARD_BG
            else:
                cell.fill.fore_color.rgb = ROW_ALT_BG

            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_HEADING if r_idx==0 else FONT_BODY
            p.font.size = Pt(9.5) if r_idx > 0 else Pt(10)
            p.font.bold = (r_idx == 0) or (c_idx == 0) or (r_idx == len(merise_data)-1)
            if r_idx == 0:
                p.font.color.rgb = TEXT_WHITE
            elif r_idx == len(merise_data) - 1 and c_idx == 2:
                p.font.color.rgb = LOGO_BRAND_GREEN
            else:
                p.font.color.rgb = TEXT_DARK

    add_footer(s17, 17)

    # ==========================================
    # SLIDE 18 : POURQUOI PU/UML ?
    # ==========================================
    s18 = prs.slides.add_slide(blank_layout)
    set_slide_background(s18, prs, LIGHT_BG)
    add_header(s18, "Justification Méthodologique", "Pourquoi Retenir le Processus Unifié et UML ?",
               "Les quatre piliers méthodologiques qui garantissent la réussite d'un projet numérique moderne")

    pu_pillars = [
        ("01", "APPROCHE ORIENTÉE OBJET", "Adéquation Technologique Parfaite",
         "Cohérence naturelle avec l'architecture en composants de React, la logique des classes JavaScript et les modèles relationnels de l'ORM Sequelize."),
        ("02", "DÉMARCHE ITÉRATIVE & INCRÉMENTALE", "Maîtrise des Risques & Agilité",
         "Développement par cycles courts permettant de tester, valider et ajuster les fonctionnalités en continu sans rupture ni effet tunnel."),
        ("03", "PROCESSUS PILOTÉ PAR LES CAS D'USAGE", "Priorité Absolue aux Besoins Métier",
         "Chaque étape de modélisation découle directement des besoins réels exprimés par les clients et les gestionnaires de SOUTARAH GROUP."),
        ("04", "STANDARD UNIVERSEL NORMALISÉ", "Communication & Documentation Fiable",
         "Le langage UML offre un cadre formel et standardisé garantissant une documentation claire et pérenne pour les futures équipes techniques.")
    ]
    coords_grid = [
        (Inches(0.8), Inches(1.85)),
        (Inches(6.8), Inches(1.85)),
        (Inches(0.8), Inches(4.35)),
        (Inches(6.8), Inches(4.35))
    ]
    for idx, (num, tag, title, desc) in enumerate(pu_pillars):
        cx, cy = coords_grid[idx]
        add_card(s18, cx, cy, Inches(5.72), Inches(2.25), top_accent_color=LOGO_BRAND_GREEN)
        add_badge(s18, cx + Inches(0.25), cy + Inches(0.2), Inches(0.55), Inches(0.3), num)
        
        tb = s18.shapes.add_textbox(cx + Inches(0.95), cy + Inches(0.2), Inches(4.5), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        pk = tf.paragraphs[0]
        pk.text = tag
        pk.font.name = FONT_HEADING
        pk.font.size = Pt(8.5)
        pk.font.bold = True
        pk.font.color.rgb = LOGO_BRAND_GREEN
        
        pt = tf.add_paragraph()
        pt.space_before = Pt(2)
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK
        
        pd = tf.add_paragraph()
        pd.space_before = Pt(4)
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_BODY

    add_footer(s18, 18)

    # ==========================================
    # SLIDE 19 : TROIS ACTEURS DU SYSTÈME
    # ==========================================
    s19 = prs.slides.add_slide(blank_layout)
    set_slide_background(s19, prs, LIGHT_BG)
    add_header(s19, "Modélisation des Acteurs", "Trois Acteurs Clés aux Rôles et Privilèges Distincts",
               "Délimitation précise des périmètres d'interaction selon le niveau d'authentification")

    acteurs = [
        ("PUBLIC", "Visiteur (Non Authentifié)",
         "Découvre SOUTARAH GROUP et ses 6 pôles d'activités.\n• Consulte le catalogue général des produits et véhicules.\n• Interagit avec l'assistant virtuel pour s'orienter.\n• Peut créer son compte client de façon autonome."),
        ("CLIENT", "Client (Authentifié)",
         "Accède à son espace personnel hautement sécurisé.\n• Configure ses dates de location et options de véhicule.\n• Gère son panier d'articles multi-pôles.\n• Génère et télécharge ses devis pro-forma officiels PDF.\n• Suit l'état d'avancement de ses réservations en direct."),
        ("ADMINISTRATION", "Administrateur (Back-Office)",
         "Supervise l'ensemble des opérations commerciales.\n• Gère la flotte automobile (ajout, entretien, tarifs).\n• Actualise le catalogue et les stocks négoce.\n• Valide ou annule les demandes de devis et réservations.\n• Analyse les KPIs d'activité et le chiffre d'affaires.")
    ]
    w_a = Inches(3.75)
    h_a = Inches(4.7)
    for idx, (tag, title, desc) in enumerate(acteurs):
        cx = Inches(0.8 + idx * 3.99)
        top_color = LOGO_LEAF_LIME if idx==2 else LOGO_BRAND_GREEN
        add_card(s19, cx, Inches(1.85), w_a, h_a, top_accent_color=top_color)
        add_badge(s19, cx + Inches(0.3), Inches(2.15), Inches(1.8), Inches(0.32), tag,
                  bg_color=LIME_LIGHT_BG,
                  txt_color=LOGO_BRAND_GREEN)
        
        tb = s19.shapes.add_textbox(cx + Inches(0.3), Inches(2.65), w_a - Inches(0.6), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = FONT_HEADING
        pt.font.size = Pt(13.5)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_DARK
        
        pd = tf.add_paragraph()
        pd.space_before = Pt(8)
        pd.text = desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_BODY

    add_footer(s19, 19)

    # ==========================================
    # SLIDE 20 : DIAGRAMME GLOBAL CAS D'USAGE (FIGURE 2)
    # ==========================================
    s20 = prs.slides.add_slide(blank_layout)
    set_slide_background(s20, prs, LIGHT_BG)
    add_header(s20, "Modélisation Fonctionnelle", "Diagramme Global des Cas d'Utilisation (Web & Mobile)",
               "Vue d'ensemble des interactions entre les acteurs et les grands packages du système")

    img_uc = "diagrammes/cas d'utilisation/cas_d'utilisation.png"
    if not os.path.exists(img_uc):
        img_uc = "extracted_pptx_images/slide_20_img_7_image.png"

    add_card(s20, Inches(0.8), Inches(1.85), Inches(7.5), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    place_image_in_box(s20, img_uc, Inches(0.95), Inches(2.0), Inches(7.2), Inches(4.3))

    # Right Explanatory Card
    add_card(s20, Inches(8.5), Inches(1.85), Inches(4.03), Inches(4.7), top_accent_color=LOGO_LEAF_LIME)
    add_badge(s20, Inches(8.75), Inches(2.1), Inches(2.8), Inches(0.32), "FIGURE 2 : SYNTHÈSE DES PACKAGES")
    
    tb_uc = s20.shapes.add_textbox(Inches(8.75), Inches(2.6), Inches(3.53), Inches(3.7))
    tf_uc = tb_uc.text_frame
    tf_uc.word_wrap = True
    tf_uc.margin_left = tf_uc.margin_top = tf_uc.margin_right = tf_uc.margin_bottom = 0

    uc_packages = [
        ("Module Authentification", "Inscription, connexion JWT, gestion du profil et déconnexion sécurisée."),
        ("Module Réservation Flotte", "Sélection de véhicule, vérification de créneau, calcul tarifaire dynamique."),
        ("Module Panier & Pro-Forma", "Agrégation multi-services, calcul des taxes et génération PDF 2 pages."),
        ("Module Back-Office Admin", "Gestion de flotte, des stocks, validation des dossiers et statistiques.")
    ]
    for i, (t, d) in enumerate(uc_packages):
        p = tf_uc.paragraphs[0] if i==0 else tf_uc.add_paragraph()
        p.space_before = Pt(8) if i>0 else Pt(0)
        p.text = f"•  {t}"
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        pd = tf_uc.add_paragraph()
        pd.space_before = Pt(2)
        pd.text = d
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED

    add_footer(s20, 20)

    # ==========================================
    # SLIDE 21 : ACTIVITÉ : COMPTE CLIENT (FIGURE 3)
    # ==========================================
    s21 = prs.slides.add_slide(blank_layout)
    set_slide_background(s21, prs, LIGHT_BG)
    add_header(s21, "Modélisation Dynamique", "Activité Clé • Création et Gestion du Compte Client",
               "Workflow d'inscription sécurisée, contrôle d'intégrité et attribution de session JWT")

    img_act1 = "diagrammes/activite/activite-compte-client.png"
    if not os.path.exists(img_act1):
        img_act1 = "extracted_pptx_images/slide_21_img_7_image.png"

    add_card(s21, Inches(0.8), Inches(1.85), Inches(5.2), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    place_image_in_box(s21, img_act1, Inches(0.95), Inches(2.0), Inches(4.9), Inches(4.3))

    add_card(s21, Inches(6.2), Inches(1.85), Inches(6.33), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s21, Inches(6.5), Inches(2.1), Inches(3.2), Inches(0.32), "FIGURE 3 : POINTS CLÉS DU TRAITEMENT")

    tb_a1 = s21.shapes.add_textbox(Inches(6.5), Inches(2.6), Inches(5.7), Inches(3.7))
    tf_a1 = tb_a1.text_frame
    tf_a1.word_wrap = True
    tf_a1.margin_left = tf_a1.margin_top = tf_a1.margin_right = tf_a1.margin_bottom = 0

    etapes_a1 = [
        ("Saisie du Formulaire d'Inscription", "L'usager fournit nom, prénom, email, téléphone et mot de passe."),
        ("Contrôle d'Intégrité & Unicité", "Vérification côté serveur que l'adresse e-mail n'est pas déjà attribuée dans MySQL."),
        ("Chiffrement Cryptographique", "Hachage sécurisé du mot de passe avec l'algorithme bcrypt (10 rounds de salage)."),
        ("Persistance & Attribution de Rôle", "Création de l'enregistrement avec le rôle 'CLIENT' par défaut dans la base."),
        ("Session JWT & Notification Mail", "Émission du token JWT signé et déclenchement immédiat d'un courriel de bienvenue via Brevo.")
    ]
    for i, (t, d) in enumerate(etapes_a1):
        p = tf_a1.paragraphs[0] if i==0 else tf_a1.add_paragraph()
        p.space_before = Pt(6) if i>0 else Pt(0)
        p.text = f"{i+1}.  {t}"
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        pd = tf_a1.add_paragraph()
        pd.space_before = Pt(2)
        pd.text = d
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED

    add_footer(s21, 21)

    # ==========================================
    # SLIDE 22 : ACTIVITÉ : RÉSERVATION VÉHICULE (FIGURE 4)
    # ==========================================
    s22 = prs.slides.add_slide(blank_layout)
    set_slide_background(s22, prs, LIGHT_BG)
    add_header(s22, "Modélisation Dynamique", "Activité Clé • Réservation d'un Véhicule Automobile",
               "Gestion des règles métier : calcul de durée, option chauffeur et contrôle de disponibilité")

    img_act2 = "diagrammes/activite/activite-reservation-vehicule.png"
    if not os.path.exists(img_act2):
        img_act2 = "extracted_pptx_images/slide_22_img_7_image.png"

    add_card(s22, Inches(0.8), Inches(1.85), Inches(6.5), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    place_image_in_box(s22, img_act2, Inches(0.95), Inches(2.0), Inches(6.2), Inches(4.3))

    add_card(s22, Inches(7.5), Inches(1.85), Inches(5.03), Inches(4.7), top_accent_color=LOGO_LEAF_LIME)
    add_badge(s22, Inches(7.75), Inches(2.1), Inches(3.2), Inches(0.32), "FIGURE 4 : RÈGLES DE GESTION")

    tb_a2 = s22.shapes.add_textbox(Inches(7.75), Inches(2.6), Inches(4.5), Inches(3.7))
    tf_a2 = tb_a2.text_frame
    tf_a2.word_wrap = True
    tf_a2.margin_left = tf_a2.margin_top = tf_a2.margin_right = tf_a2.margin_bottom = 0

    etapes_a2 = [
        ("Sélection du Véhicule", "Consultation de la fiche technique (boîte, carburant, places, tarif/jour)."),
        ("Saisie du Calendrier", "Choix des dates de départ et restitution avec calcul immédiat du nombre de jours."),
        ("Option Conducteur", "Possibilité d'ajouter l'option 'Avec chauffeur' avec majoration forfaitaire journalière."),
        ("Contrôle Conflit de Dates", "Interrogation temps réel de la base : rejet si le véhicule est déjà réservé sur la période."),
        ("Ajout au Panier Unifié", "Stockage dans la session client pour commande groupée avec devis PDF.")
    ]
    for i, (t, d) in enumerate(etapes_a2):
        p = tf_a2.paragraphs[0] if i==0 else tf_a2.add_paragraph()
        p.space_before = Pt(6) if i>0 else Pt(0)
        p.text = f"{i+1}.  {t}"
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        pd = tf_a2.add_paragraph()
        pd.space_before = Pt(2)
        pd.text = d
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED

    add_footer(s22, 22)

    # ==========================================
    # SLIDE 23 : ACTIVITÉ : ACHAT PRODUIT NÉGOCE (FIGURE 5)
    # ==========================================
    s23 = prs.slides.add_slide(blank_layout)
    set_slide_background(s23, prs, LIGHT_BG)
    add_header(s23, "Modélisation Dynamique", "Activité Clé • Achat d'un Produit du Pôle Négoce",
               "Parcours de commande des biens matériels, vérification de stock et calcul fiscal")

    img_act3 = "diagrammes/activite/achat d'un produit.png"
    if not os.path.exists(img_act3):
        img_act3 = "extracted_pptx_images/slide_23_img_7_image.png"

    add_card(s23, Inches(0.8), Inches(1.85), Inches(5.8), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    place_image_in_box(s23, img_act3, Inches(0.95), Inches(2.0), Inches(5.5), Inches(4.3))

    add_card(s23, Inches(6.8), Inches(1.85), Inches(5.73), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s23, Inches(7.05), Inches(2.1), Inches(3.2), Inches(0.32), "FIGURE 5 : FLUX OPÉRATIONNEL")

    tb_a3 = s23.shapes.add_textbox(Inches(7.05), Inches(2.6), Inches(5.2), Inches(3.7))
    tf_a3 = tb_a3.text_frame
    tf_a3.word_wrap = True
    tf_a3.margin_left = tf_a3.margin_top = tf_a3.margin_right = tf_a3.margin_bottom = 0

    etapes_a3 = [
        ("Consultation Catalogue Négoce", "Parcours des fournitures, matériels de construction et équipements industriels."),
        ("Sélection de la Quantité", "Choix du volume souhaité avec vérification du stock disponible en temps réel."),
        ("Panier Multi-Articles", "Cumul possible de plusieurs produits avec des réservations automobiles dans le même panier."),
        ("Application Règles Fiscales", "Calcul dynamique des sous-totaux HT, application de la TVA et remises éventuelles."),
        ("Validation & Devis Officiel", "Génération automatique du devis pro-forma numéroté téléchargeable instantanément.")
    ]
    for i, (t, d) in enumerate(etapes_a3):
        p = tf_a3.paragraphs[0] if i==0 else tf_a3.add_paragraph()
        p.space_before = Pt(6) if i>0 else Pt(0)
        p.text = f"{i+1}.  {t}"
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        pd = tf_a3.add_paragraph()
        pd.space_before = Pt(2)
        pd.text = d
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED

    add_footer(s23, 23)

    # ==========================================
    # SLIDE 24 : SÉQUENCE : PAIEMENT SÉCURISÉ (FIGURE 6)
    # ==========================================
    s24 = prs.slides.add_slide(blank_layout)
    set_slide_background(s24, prs, LIGHT_BG)
    add_header(s24, "Modélisation Dynamique", "Séquence • Processus de Paiement en Ligne Sécurisé",
               "Orchestration des échanges entre le Client, le Serveur Node.js et la Passerelle de paiement")

    img_seq1 = "diagrammes/sequence/sequence-paiement-en-ligne.png"
    if not os.path.exists(img_seq1):
        img_seq1 = "extracted_pptx_images/slide_24_img_7_image.png"

    add_card(s24, Inches(0.8), Inches(1.85), Inches(7.5), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    place_image_in_box(s24, img_seq1, Inches(0.95), Inches(2.0), Inches(7.2), Inches(4.3))

    add_card(s24, Inches(8.5), Inches(1.85), Inches(4.03), Inches(4.7), top_accent_color=LOGO_LEAF_LIME)
    add_badge(s24, Inches(8.75), Inches(2.1), Inches(2.5), Inches(0.32), "FIGURE 6 : SÉCURITÉ TRANSACTION")

    tb_s1 = s24.shapes.add_textbox(Inches(8.75), Inches(2.6), Inches(3.53), Inches(3.7))
    tf_s1 = tb_s1.text_frame
    tf_s1.word_wrap = True
    tf_s1.margin_left = tf_s1.margin_top = tf_s1.margin_right = tf_s1.margin_bottom = 0

    points_seq1 = [
        ("Initiation de Commande", "Le client valide son panier et choisit le mode de règlement en ligne."),
        ("Création Session Sécurisée", "L'API Node.js initialise la transaction auprès de la passerelle partenaire."),
        ("Chiffrement Bout-en-Bout", "L'usager renseigne ses informations sur un formulaire chiffré sans exposition des données."),
        ("Confirmation par Webhook", "La passerelle notifie le backend qui met à jour la commande en statut 'Payé'."),
        ("Reçu & Quittance PDF", "Génération automatique et expédition du reçu par e-mail via Brevo.")
    ]
    for i, (t, d) in enumerate(points_seq1):
        p = tf_s1.paragraphs[0] if i==0 else tf_s1.add_paragraph()
        p.space_before = Pt(6) if i>0 else Pt(0)
        p.text = f"•  {t}"
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        pd = tf_s1.add_paragraph()
        pd.space_before = Pt(2)
        pd.text = d
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED

    add_footer(s24, 24)

    # ==========================================
    # SLIDE 25 : SÉQUENCE : RÉSERVATION VÉHICULE (FIGURE 7)
    # ==========================================
    s25 = prs.slides.add_slide(blank_layout)
    set_slide_background(s25, prs, LIGHT_BG)
    add_header(s25, "Modélisation Dynamique", "Séquence • Traitement d'une Réservation de Véhicule",
               "Flux chronologique des requêtes entre l'Application Mobile, l'API REST et la Base MySQL")

    img_seq2 = "diagrammes/sequence/sequence-reservation-vehicule.png"
    if not os.path.exists(img_seq2):
        img_seq2 = "extracted_pptx_images/slide_25_img_7_image.png"

    add_card(s25, Inches(0.8), Inches(1.85), Inches(7.0), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    place_image_in_box(s25, img_seq2, Inches(0.95), Inches(2.0), Inches(6.7), Inches(4.3))

    add_card(s25, Inches(8.0), Inches(1.85), Inches(4.53), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s25, Inches(8.25), Inches(2.1), Inches(3.2), Inches(0.32), "FIGURE 7 : SYNCHRONISATION MOBILE")

    tb_s2 = s25.shapes.add_textbox(Inches(8.25), Inches(2.6), Inches(4.0), Inches(3.7))
    tf_s2 = tb_s2.text_frame
    tf_s2.word_wrap = True
    tf_s2.margin_left = tf_s2.margin_top = tf_s2.margin_right = tf_s2.margin_bottom = 0

    points_seq2 = [
        ("Soumission Mobile", "L'application mobile transmet les paramètres (dates, véhiculeId, optionChauffeur)."),
        ("Validation par Middlewares", "L'API vérifie le jeton d'authentification JWT du client demandeur."),
        ("Interrogation MySQL", "Contrôle d'absence de chevauchement sur la table 'reservations'."),
        ("Création de l'Enregistrement", "Insertion de la réservation avec statut 'En attente' et blocage du créneau."),
        ("Alerte Back-Office", "Notification instantanée transmise aux administrateurs pour traitement prioritaire.")
    ]
    for i, (t, d) in enumerate(points_seq2):
        p = tf_s2.paragraphs[0] if i==0 else tf_s2.add_paragraph()
        p.space_before = Pt(6) if i>0 else Pt(0)
        p.text = f"•  {t}"
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        pd = tf_s2.add_paragraph()
        pd.space_before = Pt(2)
        pd.text = d
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED

    add_footer(s25, 25)

    # ==========================================
    # SLIDE 26 : SÉQUENCES : ACHAT PRODUIT & INSCRIPTION (FIGURES 8 & 9)
    # ==========================================
    s26 = prs.slides.add_slide(blank_layout)
    set_slide_background(s26, prs, LIGHT_BG)
    add_header(s26, "Modélisation Dynamique", "Séquences • Inscription Utilisateur & Achat Produit",
               "Deux processus transversaux assurant l'onboarding et la concrétisation des ventes")

    img_s8 = "diagrammes/sequence/sequence-achat-produit.png"
    if not os.path.exists(img_s8):
        img_s8 = "extracted_pptx_images/slide_26_img_7_image.png"

    img_s9 = "diagrammes/sequence/sequence-authentification-inscription.png"
    if not os.path.exists(img_s9):
        img_s9 = "extracted_pptx_images/slide_26_img_9_image.png"

    # Left Box (Figure 8 : Achat Produit)
    add_card(s26, Inches(0.8), Inches(1.85), Inches(5.72), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s26, Inches(1.05), Inches(2.05), Inches(3.8), Inches(0.32), "FIGURE 8 : SÉQUENCE ACHAT PRODUIT")
    place_image_in_box(s26, img_s8, Inches(0.95), Inches(2.5), Inches(5.42), Inches(3.8))

    # Right Box (Figure 9 : Inscription Utilisateur)
    add_card(s26, Inches(6.8), Inches(1.85), Inches(5.72), Inches(4.7), top_accent_color=LOGO_LEAF_LIME)
    add_badge(s26, Inches(7.05), Inches(2.05), Inches(4.2), Inches(0.32), "FIGURE 9 : SÉQUENCE INSCRIPTION UTILISATEUR", bg_color=LIME_LIGHT_BG, txt_color=LOGO_BRAND_GREEN)
    place_image_in_box(s26, img_s9, Inches(6.95), Inches(2.5), Inches(5.42), Inches(3.8))

    add_footer(s26, 26)

    # ==========================================
    # SLIDE 27 : DIAGRAMME DE CLASSES DU SYSTÈME (FIGURE 10)
    # ==========================================
    s27 = prs.slides.add_slide(blank_layout)
    set_slide_background(s27, prs, LIGHT_BG)
    add_header(s27, "Modélisation Statique", "Structure Statique • Diagramme des Classes du Système",
               "Modélisation objet des entités métier et correspondance directe avec l'ORM Sequelize")

    img_cls = "diagrammes/classe/diagramme-de-classe.png"
    if not os.path.exists(img_cls):
        img_cls = "extracted_pptx_images/slide_27_img_7_image.png"

    add_card(s27, Inches(0.8), Inches(1.85), Inches(5.5), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    place_image_in_box(s27, img_cls, Inches(0.95), Inches(2.0), Inches(5.2), Inches(4.3))

    add_card(s27, Inches(6.5), Inches(1.85), Inches(6.03), Inches(4.7), top_accent_color=LOGO_BRAND_GREEN)
    add_badge(s27, Inches(6.75), Inches(2.1), Inches(3.4), Inches(0.32), "FIGURE 10 : STRUCTURE DES ENTITÉS")

    tb_cls = s27.shapes.add_textbox(Inches(6.75), Inches(2.6), Inches(5.5), Inches(3.7))
    tf_cls = tb_cls.text_frame
    tf_cls.word_wrap = True
    tf_cls.margin_left = tf_cls.margin_top = tf_cls.margin_right = tf_cls.margin_bottom = 0

    classes_points = [
        ("Entité User (Utilisateur)", "Attributs : id, nom, email, passwordHash, rôle (CLIENT/ADMIN), téléphone. Relation 1 à N vers Réservations et Devis."),
        ("Entité Vehicle (Véhicule)", "Attributs : id, marque, modèle, catégorie, immatriculation, tarifJour, disponibilité. Relation 1 à N vers Réservations."),
        ("Entité Reservation", "Attributs : id, dateDebut, dateFin, avecChauffeur, montantTotal, statut. Clés étrangères reliées à User et Vehicle."),
        ("Entités Product & Quote", "Gestion des articles négoce (prix, stock, pôle) et génération de devis avec items détaillés (QuoteItem)."),
        ("Mapping Sequelize ORM", "Chaque classe UML correspond rigoureusement à un modèle JavaScript Sequelize avec validation de schéma automatique.")
    ]
    for i, (t, d) in enumerate(classes_points):
        p = tf_cls.paragraphs[0] if i==0 else tf_cls.add_paragraph()
        p.space_before = Pt(6) if i>0 else Pt(0)
        p.text = f"•  {t}"
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        pd = tf_cls.add_paragraph()
        pd.space_before = Pt(2)
        pd.text = d
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED

    add_footer(s27, 27)
    print("Part 2 (Slides 16 to 27) updated with Soutarah Green theme.")
