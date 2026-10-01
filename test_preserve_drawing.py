import shutil
import docx

shutil.copyfile('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx', 'test_trans.docx')
doc = docx.Document('test_trans.docx')

def update_paragraph_text(p, new_text):
    text_runs = []
    for r in p.runs:
        if 'w:drawing' not in r._r.xml:
            text_runs.append(r)
    
    if not text_runs:
        # If no text run exists, add one
        r = p.add_run(new_text)
        return
    
    # Put translated text in first text run
    text_runs[0].text = new_text
    # Clear remaining text runs
    for r in text_runs[1:]:
        r.text = ""

update_paragraph_text(doc.paragraphs[1], "DEDICATION")
update_paragraph_text(doc.paragraphs[5], "ACKNOWLEDGEMENTS")
update_paragraph_text(doc.paragraphs[210], "2. Gantt Chart")

doc.save('test_trans.docx')

# Re-open and verify
doc2 = docx.Document('test_trans.docx')
p210 = doc2.paragraphs[210]
has_drawing = any('w:drawing' in r._r.xml for r in p210.runs)
print("Paragraph 1 text:", doc2.paragraphs[1].text)
print("Paragraph 5 text:", doc2.paragraphs[5].text)
print("Paragraph 210 text:", p210.text)
print("Paragraph 210 has drawing?", has_drawing)
