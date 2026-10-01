# -*- coding: utf-8 -*-
"""
Générateur complet de la présentation PowerPoint Soutarah Négoce Électrique
avec le fond sombre cinématique du véhicule Lynk & Co (fond_slide_dark.png)
demandé par l'utilisateur.
"""

import os
import shutil
import qrcode
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ─── COULEURS DE LA CHARTE (DARK TECH ÉLECTRIQUE) ───
COLOR_CARD_BG       = RGBColor(17, 24, 39)      # #111827 Fond de carte sombre et lisible
COLOR_CARD_BORDER   = RGBColor(55, 65, 81)      # #374151 Bordure subtile
COLOR_ACCENT_GOLD   = RGBColor(245, 158, 11)    # #F59E0B Or / Orange électrique
COLOR_ACCENT_AMBER  = RGBColor(251, 191, 36)    # #FBBF24 Jaune ambre éclatant
COLOR_ACCENT_CYAN   = RGBColor(14, 165, 233)    # #0EA5E9 Cyan technologique
COLOR_ACCENT_GREEN  = RGBColor(16, 185, 129)    # #10B981 Vert validation stock
COLOR_ACCENT_RED    = RGBColor(239, 68, 68)     # #EF4444 Rouge alerte bêtisier
COLOR_TEXT_WHITE    = RGBColor(255, 255, 255)   # #FFFFFF Blanc éclatant
COLOR_TEXT_LIGHT    = RGBColor(226, 232, 240)   # #E2E8F0 Gris clair lisible
COLOR_TEXT_MUTED    = RGBColor(156, 163, 175)   # #9CA3AF Métadonnées

FONT_TITLE = "Segoe UI"
FONT_BODY  = "Segoe UI"

LOGO_SOUTARAH   = "dist/logo-soutarah.png"
LOGO_SOGELEC    = "sogelec_transparent.png"
QR_CODE_PATH    = "qrcode_soutarah.png"
BACKGROUND_IMG  = "fond_slide_dark.png"

def ensure_assets():
    """Génère le fond sombre et le QR code si nécessaire."""
    # 1. Fond sombre cinématique
    if not os.path.exists(BACKGROUND_IMG):
        if os.path.exists('ppt/media/image1.png'):
            orig = Image.open('ppt/media/image1.png').convert('RGB')
            w, h = orig.size
            dark_bg = Image.new('RGB', (w, h), (11, 17, 32))
            blended = Image.blend(dark_bg, orig, 0.14)
            blended.save(BACKGROUND_IMG, quality=95)
    
    # 2. QR code
    if not os.path.exists(QR_CODE_PATH):
        qr = qrcode.QRCode(
            version=1,
            error_correction=qrcode.constants.ERROR_CORRECT_M,
            box_size=10,
            border=2,
        )
        qr.add_data("https://www.soutarahgroup.com")
        qr.make(fit=True)
        img = qr.make_image(fill_color="#0F172A", back_color="white")
        img.save(QR_CODE_PATH)

def set_slide_background(slide):
    """Applique l'image de fond sombre cinématique sur toute la diapositive 16:9."""
    if os.path.exists(BACKGROUND_IMG):
        slide.shapes.add_picture(BACKGROUND_IMG, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))
    
    # Ligne d'accentuation or électrique tout en haut
    top_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), Inches(13.333), Inches(0.08)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT_GOLD
    top_bar.line.fill.background()

def add_header(slide, badge_text, title_text, subtitle_text=None, badge_color=COLOR_ACCENT_GOLD):
    """En-tête avec les logos officiels Soutarah et SOGELEC."""
    # Logo Soutarah en haut à gauche
    if os.path.exists(LOGO_SOUTARAH):
        slide.shapes.add_picture(LOGO_SOUTARAH, Inches(0.8), Inches(0.3), width=Inches(1.1))

    # Logo SOGELEC en haut à droite
    if os.path.exists(LOGO_SOGELEC):
        slide.shapes.add_picture(LOGO_SOGELEC, Inches(11.2), Inches(0.32), width=Inches(1.4))

    # Badge de catégorie
    if badge_text:
        badge_box = slide.shapes.add_textbox(Inches(2.1), Inches(0.35), Inches(8.8), Inches(0.35))
        tf_b = badge_box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        run_b = p_b.add_run()
        run_b.text = badge_text.upper()
        run_b.font.name = FONT_BODY
        run_b.font.size = Pt(11)
        run_b.font.bold = True
        run_b.font.color.rgb = badge_color

    # Titre de la diapositive
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.8), Inches(0.65))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    run_t = p_t.add_run()
    run_t.text = title_text
    run_t.font.name = FONT_TITLE
    run_t.font.size = Pt(23)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_TEXT_WHITE

    # Sous-titre
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.8), Inches(0.4))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        run_s = p_s.add_run()
        run_s.text = subtitle_text
        run_s.font.name = FONT_BODY
        run_s.font.size = Pt(13)
        run_s.font.italic = True
        run_s.font.color.rgb = COLOR_TEXT_MUTED

def add_footer(slide, slide_num):
    """Pied de page discret et élégant."""
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.8), Inches(0.35))
    tf = footer_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r1 = p.add_run()
    r1.text = "SOUTARAH GROUP × SOGELEC  •  Négoce Électrique 2.0  •  soutarahgroup.com"
    r1.font.name = FONT_BODY
    r1.font.size = Pt(10)
    r1.font.color.rgb = COLOR_TEXT_MUTED
    
    r2 = p.add_run()
    r2.text = f"                                                                                           {slide_num} / 6"
    r2.font.name = FONT_BODY
    r2.font.size = Pt(10)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_ACCENT_GOLD

def create_card(slide, left, top, width, height, border_color=COLOR_CARD_BORDER, bg_color=COLOR_CARD_BG):
    """Crée une carte visuelle contrastée pour accueillir le contenu textuel."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    return card

# ─── SLIDE 1 : ACCUEIL & TON DÉCALÉ ───
def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    # Logos en en-tête
    if os.path.exists(LOGO_SOUTARAH):
        slide.shapes.add_picture(LOGO_SOUTARAH, Inches(1.0), Inches(0.6), width=Inches(1.8))
    if os.path.exists(LOGO_SOGELEC):
        slide.shapes.add_picture(LOGO_SOGELEC, Inches(10.5), Inches(0.65), width=Inches(1.8))

    # Badge Central
    badge_card = create_card(slide, Inches(3.2), Inches(0.7), Inches(6.9), Inches(0.45), border_color=COLOR_ACCENT_GOLD, bg_color=COLOR_CARD_BG)
    tf_bc = badge_card.text_frame
    tf_bc.margin_top = Inches(0.08)
    p_bc = tf_bc.paragraphs[0]
    p_bc.alignment = PP_ALIGN.CENTER
    r_bc = p_bc.add_run()
    r_bc.text = "⚡ ALLIANCE STRATÉGIQUE : SOUTARAH GROUP × SOGELEC"
    r_bc.font.name = FONT_BODY
    r_bc.font.size = Pt(12)
    r_bc.font.bold = True
    r_bc.font.color.rgb = COLOR_ACCENT_GOLD

    # Grand Titre Décalé
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(1.8))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    
    p1 = tf_t.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "Matériel électrique :\n"
    r1.font.name = FONT_TITLE
    r1.font.size = Pt(36)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_TEXT_WHITE

    r2 = p1.add_run()
    r2.text = "Et si on arrêtait de péter les plombs ?"
    r2.font.name = FONT_TITLE
    r2.font.size = Pt(44)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_ACCENT_GOLD

    # Carte Sous-Titre
    sub_card = create_card(slide, Inches(1.0), Inches(3.6), Inches(11.333), Inches(1.2), border_color=COLOR_CARD_BORDER, bg_color=COLOR_CARD_BG)
    tf_sub = sub_card.text_frame
    tf_sub.margin_top = Inches(0.18)
    tf_sub.margin_left = Inches(0.4)
    tf_sub.margin_right = Inches(0.4)
    
    p_sub = tf_sub.paragraphs[0]
    p_sub.alignment = PP_ALIGN.CENTER
    r_sub = p_sub.add_run()
    r_sub.text = "Devis instantanés, stock visible en live & livraison directe sur chantier par Soutarah Group."
    r_sub.font.name = FONT_TITLE
    r_sub.font.size = Pt(18)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_ACCENT_AMBER

    p_sub2 = tf_sub.add_paragraph()
    p_sub2.alignment = PP_ALIGN.CENTER
    r_sub2 = p_sub2.add_run()
    r_sub2.text = "⚡ Chiffrage immédiat 24/7   •   📦 Visibilité stock en temps réel   •   🚚 Acheminement direct sur site"
    r_sub2.font.name = FONT_BODY
    r_sub2.font.size = Pt(13)
    r_sub2.font.color.rgb = COLOR_TEXT_LIGHT

    # Carte Accroche Punchline
    acc_card = create_card(slide, Inches(1.0), Inches(5.1), Inches(11.333), Inches(1.5), border_color=COLOR_ACCENT_CYAN, bg_color=RGBColor(16, 26, 46))
    tf_acc = acc_card.text_frame
    tf_acc.margin_top = Inches(0.25)
    tf_acc.margin_left = Inches(0.5)
    tf_acc.margin_right = Inches(0.5)

    p_a1 = tf_acc.paragraphs[0]
    p_a1.alignment = PP_ALIGN.CENTER
    r_a1 = p_a1.add_run()
    r_a1.text = "« Parce qu’un technicien qui attend du matériel, c’est de l’argent qui s'envole en fumée. »"
    r_a1.font.name = FONT_TITLE
    r_a1.font.size = Pt(17)
    r_a1.font.italic = True
    r_a1.font.bold = True
    r_a1.font.color.rgb = COLOR_TEXT_WHITE

    p_a2 = tf_acc.add_paragraph()
    p_a2.alignment = PP_ALIGN.CENTER
    r_a2 = p_a2.add_run()
    r_a2.text = "👉 Découvrez le négoce version 2.0 avec Soutarah Group."
    r_a2.font.name = FONT_BODY
    r_a2.font.size = Pt(15)
    r_a2.font.bold = True
    r_a2.font.color.rgb = COLOR_ACCENT_GOLD

    add_footer(slide, 1)

# ─── SLIDE 2 : LE CONSTAT (LE BÊTISIER) ───
def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    add_header(
        slide,
        badge_text="⚠️ LE CONSTAT CHANTIER",
        title_text="Scène classique sur chantier...",
        subtitle_text="(Toute ressemblance avec votre quotidien est purement volontaire)",
        badge_color=COLOR_ACCENT_RED
    )

    cards_data = [
        {
            "icon": "⏳",
            "title": "Le devis \"fantôme\"",
            "quote": "« Demandé lundi, relancé mercredi, reçu... le mois prochain. »",
            "impact": "Résultat : Chantier bloqué, retards de livraison et plannings explosés.",
            "color": COLOR_ACCENT_RED
        },
        {
            "icon": "📦",
            "title": "La boîte vide",
            "quote": "« Faire 2 heures d'embouteillages pour entendre : \"Ah mon frère, c'est fini en stock !\" »",
            "impact": "Résultat : Temps perdu dans les embouteillages d'Abidjan et carburant gaspillé.",
            "color": COLOR_ACCENT_GOLD
        },
        {
            "icon": "⚡",
            "title": "Le mauvais composant",
            "quote": "« Le disjoncteur reçu n'est pas le bon, mais \"ça s'adapte avec un bout de scotch\" (spoilers : non). »",
            "impact": "Résultat : Risques graves d'incendie, non-conformité et reprise coûteuse.",
            "color": COLOR_ACCENT_RED
        },
        {
            "icon": "🚚",
            "title": "Le marathon logistique",
            "quote": "« Devoir transformer le chef de chantier en livreur de câbles improvisé. »",
            "impact": "Résultat : Main-d'œuvre qualifiée détournée de sa vraie mission sur le chantier.",
            "color": COLOR_ACCENT_GOLD
        }
    ]

    positions = [
        (Inches(0.8), Inches(2.0)),
        (Inches(6.8), Inches(2.0)),
        (Inches(0.8), Inches(4.5)),
        (Inches(6.8), Inches(4.5)),
    ]

    w_card = Inches(5.7)
    h_card = Inches(2.3)

    for i, data in enumerate(cards_data):
        pos_x, pos_y = positions[i]
        card = create_card(slide, pos_x, pos_y, w_card, h_card, border_color=data["color"], bg_color=COLOR_CARD_BG)
        tf = card.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.2)
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = f"{data['icon']}  {data['title']}"
        r1.font.name = FONT_TITLE
        r1.font.size = Pt(17)
        r1.font.bold = True
        r1.font.color.rgb = data["color"]

        p2 = tf.add_paragraph()
        p2.space_before = Pt(6)
        r2 = p2.add_run()
        r2.text = data["quote"]
        r2.font.name = FONT_BODY
        r2.font.size = Pt(13)
        r2.font.italic = True
        r2.font.color.rgb = COLOR_TEXT_WHITE

        p3 = tf.add_paragraph()
        p3.space_before = Pt(8)
        r3 = p3.add_run()
        r3.text = data["impact"]
        r3.font.name = FONT_BODY
        r3.font.size = Pt(12)
        r3.font.color.rgb = COLOR_TEXT_MUTED

    add_footer(slide, 2)

# ─── SLIDE 3 : LA SOLUTION (SUPER-POUVOIRS) ───
def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    add_header(
        slide,
        badge_text="💡 LA RÉVOLUTION SOUTARAH NÉGOCE",
        title_text="Vous avez le chantier, nous avons les étincelles !",
        subtitle_text="Les 4 super-pouvoirs de notre plateforme pour un approvisionnement sans prise de tête",
        badge_color=COLOR_ACCENT_CYAN
    )

    superpowers = [
        {
            "icon": "⚡",
            "title": "Devis plus rapide\nque l'éclair",
            "desc": "Chiffrez vos besoins en ligne en temps réel.",
            "detail": "Accédez instantanément aux prix officiels sans attendre des jours. Émission de devis pro-forma certifié en 1 clic.",
            "color": COLOR_ACCENT_GOLD
        },
        {
            "icon": "👁️",
            "title": "Transparence\n\"Rayons X\"",
            "desc": "Vous voyez le stock exact disponible avant de cliquer.",
            "detail": "Zéro fausse promesse. Si c'est affiché disponible sur la plateforme, le produit est physiquement en rayon.",
            "color": COLOR_ACCENT_CYAN
        },
        {
            "icon": "🤝",
            "title": "Écosystème\nde confiance",
            "desc": "Soutenu par des leaders comme SOGELEC.",
            "detail": "Partenariat direct garantissant du matériel 100% certifié constructeur, normé et d'origine contrôlée.",
            "color": COLOR_ACCENT_GREEN
        },
        {
            "icon": "🚚",
            "title": "Livraison\n\"Zéro déplacement\"",
            "desc": "On vous livre directement sur votre chantier.",
            "detail": "Acheminement express sur site à Abidjan et partout à l'intérieur du pays. Vos équipes restent concentrées.",
            "color": COLOR_ACCENT_AMBER
        }
    ]

    card_w = Inches(2.75)
    card_h = Inches(4.7)
    gap = Inches(0.25)
    start_x = Inches(0.8)
    top_y = Inches(2.0)

    for i, sp in enumerate(superpowers):
        pos_x = start_x + i * (card_w + gap)
        card = create_card(slide, pos_x, top_y, card_w, card_h, border_color=sp["color"], bg_color=COLOR_CARD_BG)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.25)
        tf.word_wrap = True

        p_icon = tf.paragraphs[0]
        p_icon.alignment = PP_ALIGN.CENTER
        r_icon = p_icon.add_run()
        r_icon.text = sp["icon"]
        r_icon.font.size = Pt(36)

        p_tit = tf.add_paragraph()
        p_tit.alignment = PP_ALIGN.CENTER
        p_tit.space_before = Pt(8)
        r_tit = p_tit.add_run()
        r_tit.text = sp["title"]
        r_tit.font.name = FONT_TITLE
        r_tit.font.size = Pt(16)
        r_tit.font.bold = True
        r_tit.font.color.rgb = sp["color"]

        p_desc = tf.add_paragraph()
        p_desc.alignment = PP_ALIGN.CENTER
        p_desc.space_before = Pt(12)
        r_desc = p_desc.add_run()
        r_desc.text = sp["desc"]
        r_desc.font.name = FONT_BODY
        r_desc.font.size = Pt(13)
        r_desc.font.bold = True
        r_desc.font.color.rgb = COLOR_TEXT_WHITE

        p_sep = tf.add_paragraph()
        p_sep.alignment = PP_ALIGN.CENTER
        p_sep.space_before = Pt(8)
        r_sep = p_sep.add_run()
        r_sep.text = "─────────"
        r_sep.font.size = Pt(10)
        r_sep.font.color.rgb = COLOR_CARD_BORDER

        p_det = tf.add_paragraph()
        p_det.space_before = Pt(8)
        r_det = p_det.add_run()
        r_det.text = sp["detail"]
        r_det.font.name = FONT_BODY
        r_det.font.size = Pt(11)
        r_det.font.color.rgb = COLOR_TEXT_LIGHT

    add_footer(slide, 3)

# ─── SLIDE 4 : COMMENT ÇA MARCHE (4 ÉTAPES) ───
def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    add_header(
        slide,
        badge_text="⚡ PARCOURS CLIENT FLUIDE & RAPIDE",
        title_text="Moins de blabla, plus de watts.",
        subtitle_text="Votre matériel électrique livré sur chantier en 4 étapes chrono",
        badge_color=COLOR_ACCENT_GOLD
    )

    steps = [
        {
            "num": "01",
            "action": "CLIQUEZ",
            "icon": "🖱️",
            "subtitle": "Catalogue en ligne 24/7",
            "text": "Trouvez votre câble, disjoncteur ou armoire sur soutarahgroup.com en quelques clics.",
            "color": COLOR_ACCENT_CYAN
        },
        {
            "num": "02",
            "action": "VÉRIFIEZ",
            "icon": "🟢",
            "subtitle": "Disponibilité immédiate",
            "text": "Le stock est au vert ? C’est réservé instantanément pour vous dans nos entrepôts !",
            "color": COLOR_ACCENT_GREEN
        },
        {
            "num": "03",
            "action": "CHIFFREZ & VALIDEZ",
            "icon": "📑",
            "subtitle": "Devis express en direct",
            "text": "Devis instantané, prix fermes et validation immédiate de votre commande sans attente.",
            "color": COLOR_ACCENT_GOLD
        },
        {
            "num": "04",
            "action": "LIVRAISON DIRECTE",
            "icon": "🚚",
            "subtitle": "Déchargement sur site",
            "text": "On décharge directement sur votre chantier. Gardez les mains sur vos outils !",
            "color": COLOR_ACCENT_AMBER
        }
    ]

    card_w = Inches(2.75)
    card_h = Inches(4.7)
    gap = Inches(0.25)
    start_x = Inches(0.8)
    top_y = Inches(2.0)

    for i, st in enumerate(steps):
        pos_x = start_x + i * (card_w + gap)
        card = create_card(slide, pos_x, top_y, card_w, card_h, border_color=st["color"], bg_color=COLOR_CARD_BG)
        tf = card.text_frame
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.3)
        tf.word_wrap = True

        p_num = tf.paragraphs[0]
        r_num = p_num.add_run()
        r_num.text = st["num"]
        r_num.font.name = FONT_TITLE
        r_num.font.size = Pt(38)
        r_num.font.bold = True
        r_num.font.color.rgb = st["color"]

        p_act = tf.add_paragraph()
        p_act.space_before = Pt(4)
        r_act = p_act.add_run()
        r_act.text = f"{st['icon']}  {st['action']}"
        r_act.font.name = FONT_TITLE
        r_act.font.size = Pt(17)
        r_act.font.bold = True
        r_act.font.color.rgb = COLOR_TEXT_WHITE

        p_sub = tf.add_paragraph()
        p_sub.space_before = Pt(4)
        r_sub = p_sub.add_run()
        r_sub.text = st["subtitle"]
        r_sub.font.name = FONT_BODY
        r_sub.font.size = Pt(12)
        r_sub.font.bold = True
        r_sub.font.color.rgb = st["color"]

        p_sep = tf.add_paragraph()
        p_sep.space_before = Pt(10)
        r_sep = p_sep.add_run()
        r_sep.text = "─────────"
        r_sep.font.size = Pt(10)
        r_sep.font.color.rgb = COLOR_CARD_BORDER

        p_txt = tf.add_paragraph()
        p_txt.space_before = Pt(10)
        r_txt = p_txt.add_run()
        r_txt.text = st["text"]
        r_txt.font.name = FONT_BODY
        r_txt.font.size = Pt(12)
        r_txt.font.color.rgb = COLOR_TEXT_LIGHT

    add_footer(slide, 4)

# ─── SLIDE 5 : POURQUOI SOUTARAH (ALLIANCE & RÉSEAU) ───
def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    add_header(
        slide,
        badge_text="🛡️ AVANTAGES CONCURRENTIELS DÉCISIFS",
        title_text="La puissance d'un grand réseau, la réactivité d'un partenaire local.",
        subtitle_text="4 arguments chocs pour sécuriser l'approvisionnement de vos chantiers",
        badge_color=COLOR_ACCENT_AMBER
    )

    arguments = [
        {
            "icon": "🏆",
            "title": "Réseau Partenaires Solide",
            "quote": "En collaboration directe avec des géants comme SOGELEC",
            "desc": "Bénéficiez d'une chaîne logistique sans faille, de stocks massifs réservés et des meilleurs tarifs négociés du marché ivoirien.",
            "color": COLOR_ACCENT_GOLD
        },
        {
            "icon": "⏱️",
            "title": "Zéro Temps Mort",
            "quote": "Vos équipes travaillent, elles n'attendent plus après un rouleau de câble",
            "desc": "Chaque minute de retard coûte cher. Nos livraisons synchronisées garantissent la continuité continue de vos chantiers.",
            "color": COLOR_ACCENT_CYAN
        },
        {
            "icon": "🛡️",
            "title": "Matériel 100% Conforme",
            "quote": "Pas de contrefaçon, pas de bricolage, de la sécurité certifiée",
            "desc": "Tous nos composants respectent scrupuleusement les normes de sécurité en vigueur. Zéro risque pour vos installations.",
            "color": COLOR_ACCENT_GREEN
        },
        {
            "icon": "📞",
            "title": "Un SAV Toujours Branché",
            "quote": "Une équipe réactive au téléphone et sur le terrain",
            "desc": "Un doute technique ? Une référence spécifique ? Nos ingénieurs et conseillers dédiés vous accompagnent à chaque étape.",
            "color": COLOR_ACCENT_AMBER
        }
    ]

    positions = [
        (Inches(0.8), Inches(2.0)),
        (Inches(6.8), Inches(2.0)),
        (Inches(0.8), Inches(4.5)),
        (Inches(6.8), Inches(4.5)),
    ]

    w_card = Inches(5.7)
    h_card = Inches(2.3)

    for i, arg in enumerate(arguments):
        pos_x, pos_y = positions[i]
        card = create_card(slide, pos_x, pos_y, w_card, h_card, border_color=arg["color"], bg_color=COLOR_CARD_BG)
        tf = card.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.2)
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = f"{arg['icon']}  {arg['title']}"
        r1.font.name = FONT_TITLE
        r1.font.size = Pt(17)
        r1.font.bold = True
        r1.font.color.rgb = arg["color"]

        p2 = tf.add_paragraph()
        p2.space_before = Pt(6)
        r2 = p2.add_run()
        r2.text = f"« {arg['quote']} »"
        r2.font.name = FONT_BODY
        r2.font.size = Pt(12)
        r2.font.bold = True
        r2.font.color.rgb = COLOR_TEXT_WHITE

        p3 = tf.add_paragraph()
        p3.space_before = Pt(6)
        r3 = p3.add_run()
        r3.text = arg["desc"]
        r3.font.name = FONT_BODY
        r3.font.size = Pt(11)
        r3.font.color.rgb = COLOR_TEXT_LIGHT

    add_footer(slide, 5)

# ─── SLIDE 6 : CONCLUSION & APPEL À L'ACTION ───
def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    add_header(
        slide,
        badge_text="🚀 PASSEZ À L'ACTION DÈS AUJOURD'HUI",
        title_text="Branchez vos chantiers sur le bon courant !",
        subtitle_text="Optimisez vos approvisionnements électriques avec Soutarah Group & SOGELEC",
        badge_color=COLOR_ACCENT_GOLD
    )

    # Colonne gauche : Coordonnées
    left_card = create_card(slide, Inches(0.8), Inches(2.0), Inches(7.5), Inches(4.7), border_color=COLOR_ACCENT_GOLD, bg_color=COLOR_CARD_BG)
    tf_l = left_card.text_frame
    tf_l.margin_left = Inches(0.4)
    tf_l.margin_right = Inches(0.4)
    tf_l.margin_top = Inches(0.35)
    tf_l.word_wrap = True

    p_c1 = tf_l.paragraphs[0]
    r_c1 = p_c1.add_run()
    r_c1.text = "🌐  Plateforme Web Officielle :"
    r_c1.font.name = FONT_TITLE
    r_c1.font.size = Pt(14)
    r_c1.font.color.rgb = COLOR_TEXT_MUTED

    p_c2 = tf_l.add_paragraph()
    p_c2.space_before = Pt(2)
    r_c2 = p_c2.add_run()
    r_c2.text = "www.soutarahgroup.com"
    r_c2.font.name = FONT_TITLE
    r_c2.font.size = Pt(22)
    r_c2.font.bold = True
    r_c2.font.color.rgb = COLOR_ACCENT_GOLD

    p_c3 = tf_l.add_paragraph()
    p_c3.space_before = Pt(14)
    r_c3 = p_c3.add_run()
    r_c3.text = "⚡  Section Dédiée : "
    r_c3.font.name = FONT_BODY
    r_c3.font.size = Pt(14)
    r_c3.font.color.rgb = COLOR_TEXT_MUTED
    r_c3b = p_c3.add_run()
    r_c3b.text = "Négoce Électrique B2B"
    r_c3b.font.bold = True
    r_c3b.font.color.rgb = COLOR_TEXT_WHITE

    p_c4 = tf_l.add_paragraph()
    p_c4.space_before = Pt(6)
    r_c4 = p_c4.add_run()
    r_c4.text = "🤝  Espace Partenaires : "
    r_c4.font.name = FONT_BODY
    r_c4.font.size = Pt(14)
    r_c4.font.color.rgb = COLOR_TEXT_MUTED
    r_c4b = p_c4.add_run()
    r_c4b.text = "SOGELEC & Soutarah Group"
    r_c4b.font.bold = True
    r_c4b.font.color.rgb = COLOR_TEXT_WHITE

    p_c5 = tf_l.add_paragraph()
    p_c5.space_before = Pt(6)
    r_c5 = p_c5.add_run()
    r_c5.text = "📍  Couverture : "
    r_c5.font.name = FONT_BODY
    r_c5.font.size = Pt(14)
    r_c5.font.color.rgb = COLOR_TEXT_MUTED
    r_c5b = p_c5.add_run()
    r_c5b.text = "Abidjan et tout le territoire ivoirien"
    r_c5b.font.bold = True
    r_c5b.font.color.rgb = COLOR_TEXT_WHITE

    p_sig_box = tf_l.add_paragraph()
    p_sig_box.space_before = Pt(24)
    r_s1 = p_sig_box.add_run()
    r_s1.text = "SOUTARAH GROUP\n"
    r_s1.font.name = FONT_TITLE
    r_s1.font.size = Pt(16)
    r_s1.font.bold = True
    r_s1.font.color.rgb = COLOR_ACCENT_AMBER

    r_s2 = p_sig_box.add_run()
    r_s2.text = "Le réflexe matériel qui vous évite le court-circuit."
    r_s2.font.name = FONT_BODY
    r_s2.font.size = Pt(14)
    r_s2.font.italic = True
    r_s2.font.bold = True
    r_s2.font.color.rgb = COLOR_TEXT_WHITE

    # Colonne droite : QR Code
    right_card = create_card(slide, Inches(8.6), Inches(2.0), Inches(3.933), Inches(4.7), border_color=COLOR_ACCENT_CYAN, bg_color=COLOR_CARD_BG)
    tf_r = right_card.text_frame
    tf_r.margin_top = Inches(0.2)
    tf_r.word_wrap = True

    p_qr_t = tf_r.paragraphs[0]
    p_qr_t.alignment = PP_ALIGN.CENTER
    r_qr_t = p_qr_t.add_run()
    r_qr_t.text = "📱 COMMANDEZ EN LIGNE"
    r_qr_t.font.name = FONT_TITLE
    r_qr_t.font.size = Pt(13)
    r_qr_t.font.bold = True
    r_qr_t.font.color.rgb = COLOR_ACCENT_CYAN

    # Image du QR Code
    if os.path.exists(QR_CODE_PATH):
        slide.shapes.add_picture(QR_CODE_PATH, Inches(9.35), Inches(2.5), width=Inches(2.4), height=Inches(2.4))

    # Légende parfaitement positionnée sous le QR Code
    sub_box = slide.shapes.add_textbox(Inches(8.7), Inches(5.1), Inches(3.733), Inches(1.3))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_qr_sub = tf_sub.paragraphs[0]
    p_qr_sub.alignment = PP_ALIGN.CENTER
    r_qr_sub = p_qr_sub.add_run()
    r_qr_sub.text = "« Scannez avant la prochaine\ncoupure de courant ! »"
    r_qr_sub.font.name = FONT_BODY
    r_qr_sub.font.size = Pt(12)
    r_qr_sub.font.bold = True
    r_qr_sub.font.italic = True
    r_qr_sub.font.color.rgb = COLOR_ACCENT_GOLD

    add_footer(slide, 6)

def main():
    print("=" * 70)
    print("  GÉNÉRATION AVEC LE FOND SOMBRE CINÉMATIQUE (Lynk & Co)")
    print("=" * 70)

    # 1. Vérification / génération des assets
    ensure_assets()
    print(f"[1/4] Image de fond prête : {BACKGROUND_IMG}")

    # 2. Initialisation 16:9
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 3. Construction des 6 slides
    print("[2/4] Création des 6 diapositives avec le fond sombre...")
    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    print("      -> 6 diapositives créées avec succès.")

    # 4. Enregistrement
    out_files = [
        "Presentation_Soutarah_Negoce_Electrique.pptx",
        "Présentation1_Soutarah_Negoce.pptx"
    ]
    for out in out_files:
        prs.save(out)
        print(f"[3/4] Enregistré : '{out}' ({os.path.getsize(out):,} octets)")

    # Tenter également d'écraser Présentation1.pptx si déverrouillé
    try:
        prs.save("Présentation1.pptx")
        print("      -> 'Présentation1.pptx' mis à jour directement avec succès !")
    except PermissionError:
        print("      -> Note : 'Présentation1.pptx' est actuellement verrouillé par PowerPoint.")
        print("         Utilisez 'Présentation1_Soutarah_Negoce.pptx' ou 'Presentation_Soutarah_Negoce_Electrique.pptx'.")

    print("\n" + "=" * 70)
    print("  SUCCÈS : LE POWERPOINT EST PRÊT AVEC LE FOND SOMBRE !")
    print("=" * 70)

if __name__ == "__main__":
    main()
