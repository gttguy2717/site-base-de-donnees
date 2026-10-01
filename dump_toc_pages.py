import sys
from docx import Document

doc = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

print("=== TABLE DES MATIERES (FR) ===")
for i in range(714, len(doc.paragraphs)):
    t = doc.paragraphs[i].text.strip()
    if t:
        clean = t.encode('ascii', 'replace').decode('ascii')
        print(f"[{i}] {clean}")
