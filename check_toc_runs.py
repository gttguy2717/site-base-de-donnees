from docx import Document

doc = Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

print("=== P[30] runs (Sommaire item) ===")
p30 = doc.paragraphs[30]
for j, r in enumerate(p30.runs):
    print(f"  run[{j}]: text={repr(r.text)}, bold={r.bold}")

print("\n=== P[55] runs (Figure item) ===")
p55 = doc.paragraphs[55]
for j, r in enumerate(p55.runs):
    print(f"  run[{j}]: text={repr(r.text)}, bold={r.bold}")

print("\n=== P[87] runs (Table item) ===")
p87 = doc.paragraphs[87]
for j, r in enumerate(p87.runs):
    print(f"  run[{j}]: text={repr(r.text)}, bold={r.bold}")
