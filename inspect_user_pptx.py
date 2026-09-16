import pptx
import sys
sys.stdout.reconfigure(encoding='utf-8')

file_path = 'Soutenance stage plateforme numerique Soutarah Group.pptx'
prs = pptx.Presentation(file_path)

print(f"Total slides: {len(prs.slides)}")
print(f"Slide dimensions: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f}")

for idx in range(min(10, len(prs.slides))):
    slide = prs.slides[idx]
    print(f"\n================ SLIDE {idx+1} ================")
    for s in slide.shapes:
        color_fill = ""
        if hasattr(s, 'fill') and s.fill.type == 1:
            try:
                color_fill = f"fill={s.fill.fore_color.rgb}"
            except:
                pass
        line_col = ""
        if hasattr(s, 'line') and s.line.fill.type == 1:
            try:
                line_col = f"line={s.line.color.rgb}"
            except:
                pass
        
        text = ""
        font_info = ""
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if p.text.strip():
                    text = p.text.strip()[:40]
                    if p.runs:
                        r = p.runs[0]
                        f_name = r.font.name or ""
                        f_size = f"{r.font.size.pt:.1f}pt" if r.font.size else ""
                        try:
                            f_col = str(r.font.color.rgb)
                        except:
                            f_col = ""
                        font_info = f"[{f_name} {f_size} {f_col}]"
                    break
        print(f"  {s.name} ({s.shape_type}) pos=({s.left/914400:.2f}, {s.top/914400:.2f}, w={s.width/914400:.2f}, h={s.height/914400:.2f}) {color_fill} {line_col} {font_info} text: {text}")
