import docx
import json

doc = docx.Document('Rapport_de_Stage_SORHO_DAVID_SOUTARAH_PAGES_CORRIGEES.docx')

data = []

# Paragraphs
for i, p in enumerate(doc.paragraphs):
    has_drawing = any('w:drawing' in r._r.xml for r in p.runs)
    text = p.text
    if text.strip():
        data.append({
            'type': 'p',
            'idx': i,
            'style': p.style.name,
            'has_drawing': has_drawing,
            'text': text
        })

# Tables
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            for p_idx, p in enumerate(cell.paragraphs):
                text = p.text
                if text.strip():
                    data.append({
                        'type': 'cell',
                        'table_idx': t_idx,
                        'row_idx': r_idx,
                        'col_idx': c_idx,
                        'p_idx': p_idx,
                        'text': text
                    })

print(f"Total non-empty text elements: {len(data)}")
with open('doc_content_fr.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print("Saved doc_content_fr.json successfully!")
